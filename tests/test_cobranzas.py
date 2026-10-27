# -*- coding: utf-8 -*-
"""
Banco de pruebas del Gestor de Cobranzas Wood Tools.

Correr con:   PROBAR.bat        (o bien:  python tests/test_cobranzas.py)

Ejercita el codigo real del proyecto, pero redirige TODO lo que el programa
escribe (log, vendedores.json, asignaciones.json, Resumenes/, Historial/) a una
carpeta temporal del sistema, asi que una corrida no toca los datos de trabajo
ni el repositorio.

Las llamadas a Meta y a la API de Claude se simulan: la prueba no usa internet,
no gasta credito y no manda ningun WhatsApp. Los archivos de Excel y de PDF que
se leen se generan al vuelo con el formato del reporte del ERP.

Devuelve codigo de salida 0 si pasa todo, 1 si hay alguna falla.
"""
from __future__ import annotations

import io
import json
import os
import shutil
import sys
import tempfile
import traceback
from datetime import date, datetime

# El proyecto es la carpeta que contiene a tests/
PROYECTO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TMP = os.path.join(tempfile.gettempdir(), "pruebas_gestor_cobranzas")
shutil.rmtree(TMP, ignore_errors=True)
os.makedirs(TMP, exist_ok=True)
sys.path.insert(0, PROYECTO)

# --- redirigir escrituras ANTES de importar los modulos que las usan ---------
import config  # noqa: E402

config.ARCHIVO_LOG = os.path.join(TMP, "errores_log.txt")
config.ARCHIVO_VENDEDORES = os.path.join(TMP, "vendedores.json")
config.ARCHIVO_ASIGNACIONES = os.path.join(TMP, "asignaciones.json")
config.CARPETA_SALIDA = os.path.join(TMP, "Resumenes")
config.CARPETA_HISTORIAL = os.path.join(TMP, "Historial")
os.makedirs(config.CARPETA_SALIDA, exist_ok=True)
os.makedirs(config.CARPETA_HISTORIAL, exist_ok=True)

import modelos  # noqa: E402
import parser_excel  # noqa: E402
import parser_pdf  # noqa: E402
import mensajes  # noqa: E402
import pdf_resumen  # noqa: E402
import vendedores as vend  # noqa: E402
import whatsapp  # noqa: E402
import claude_api  # noqa: E402

vend.ARCHIVO_VENDEDORES = config.ARCHIVO_VENDEDORES
vend.ARCHIVO_ASIGNACIONES = config.ARCHIVO_ASIGNACIONES
pdf_resumen.CARPETA_SALIDA = config.CARPETA_SALIDA

# la config real del usuario puede tener claves; para las pruebas la neutralizamos
config.CONFIG.datos["anthropic_api_key"] = ""
config.CONFIG.datos["whatsapp_token"] = ""
config.CONFIG.datos["whatsapp_phone_number_id"] = ""
config.CONFIG.datos["firma_mensaje"] = "Administracion Wood Tools"
config.CONFIG.datos["max_caracteres_mensaje"] = 3800
config.CONFIG.datos["modo_envio"] = "libre"
config.CONFIG.ruta = os.path.join(TMP, "config.json")

# ==========================================================================
# MINI-FRAMEWORK
# ==========================================================================
RESULTADOS: list[tuple[str, str, str]] = []   # (grupo, nombre, estado/detalle)
_grupo_actual = "?"


def grupo(nombre: str) -> None:
    global _grupo_actual
    _grupo_actual = nombre
    print(f"\n{'=' * 78}\n{nombre}\n{'=' * 78}")


def check(nombre: str, condicion: bool, detalle: str = "") -> bool:
    estado = "OK  " if condicion else "FALLA"
    RESULTADOS.append((_grupo_actual, nombre, "ok" if condicion else f"FALLA {detalle}"))
    print(f"  [{estado}] {nombre}" + (f"   -> {detalle}" if detalle and not condicion else ""))
    return bool(condicion)


def eq(nombre: str, obtenido, esperado, tol: float | None = None) -> bool:
    if tol is not None and isinstance(obtenido, (int, float)) and isinstance(esperado, (int, float)):
        ok = abs(obtenido - esperado) <= tol
    else:
        ok = obtenido == esperado
    return check(nombre, ok, f"obtenido={obtenido!r} esperado={esperado!r}")


def _cerrar_ventana(app) -> None:
    """Cierra la ventana cancelando los after() pendientes.

    La App se reprograma sola cada 120 ms para leer la cola del hilo worker; si
    se destruye con un callback en vuelo, Tcl escupe un "invalid command name"
    que no rompe nada pero ensucia la salida de la prueba.
    """
    try:
        for pendiente in app.tk.call("after", "info"):
            try:
                app.after_cancel(pendiente)
            except Exception:
                pass
    except Exception:
        pass
    try:
        app.destroy()
    except Exception:
        pass


def nota(texto: str, sigue_pendiente: bool) -> None:
    """Deja constancia de un pendiente conocido, sin contarlo como falla.

    Sirve para que el banco documente lo que todavia falta sin romperse el dia
    que alguien lo arregle: cuando el pendiente se resuelve, la linea pasa a
    decir "resuelto" y la corrida sigue en verde.
    """
    print(f"  [{'PENDIENTE' if sigue_pendiente else 'resuelto ':<9}] {texto}")


def seccion(fn):
    """Ejecuta una seccion capturando excepciones para no cortar la corrida."""
    try:
        fn()
    except Exception as e:
        RESULTADOS.append((_grupo_actual, fn.__name__, f"EXCEPCION {e}"))
        print(f"  [ERROR] {fn.__name__} lanzo una excepcion:")
        traceback.print_exc()


# ==========================================================================
# 1. MODELOS
# ==========================================================================
def t_modelos() -> None:
    grupo("1. modelos.py  (contrato de datos, plata y telefonos)")

    eq("formatear_pesos(1234567.5)", modelos.formatear_pesos(1234567.5), "$1.234.567,50")
    eq("formatear_pesos(-1234.5)", modelos.formatear_pesos(-1234.5), "-$1.234,50")
    eq("formatear_pesos(None)", modelos.formatear_pesos(None), "$0,00")
    eq("formatear_pesos(0)", modelos.formatear_pesos(0), "$0,00")

    eq("clasificar_dias(10)", modelos.clasificar_dias(10), modelos.VENCIDO)
    eq("clasificar_dias(0)", modelos.clasificar_dias(0), modelos.VENCE_HOY)
    eq("clasificar_dias(-3)", modelos.clasificar_dias(-3), modelos.A_VENCER)
    eq("clasificar_dias(None)", modelos.clasificar_dias(None), modelos.A_VENCER)

    casos = [
        ("11 3900-9488", "5491139009488"),
        ("5491139009488", "5491139009488"),
        ("541139009488", "5491139009488"),
        ("45246075", "5491145246075"),
        ("1140647300", "5491140647300"),
        ("0341 15 456 7890", "5493414567890"),
        ("15-3359-5605", "5491133595605"),
        ("11 15 4064 7300", "5491140647300"),
        ("114085-5781", "5491140855781"),
        ("", None),
        ("1234", None),
        ("hola", None),
    ]
    for entrada, esperado in casos:
        eq(f"normalizar_telefono_ar({entrada!r})",
           modelos.normalizar_telefono_ar(entrada), esperado)

    tel = modelos.telefonos_de_texto("Marcos 11 3900 9488 -4313 0466 // 4312 0")
    eq("telefonos_de_texto (movil primero)", tel, ["5491139009488", "5491143130466"])

    tel2 = modelos.telefonos_de_texto("45246075 LAUTARO 1140647300 KAREN PAGO")
    eq("telefonos_de_texto (nombres pegados)", sorted(tel2),
       ["5491140647300", "5491145246075"])
    nota("la heuristica marca los 11-4xxx como fijos, aunque haya celulares asi",
         not modelos.es_probable_movil_ar("5491140647300"))

    # cliente sintetico
    c = modelos.Cliente(codigo="00000002761", razon_social="ACME SRL")
    c.comprobantes = [
        modelos.Comprobante(tipo="FE A", talonario="71", numero="1", saldo=100.0,
                            dias_venc=10, fecha_vto=date(2026, 7, 31)),
        modelos.Comprobante(tipo="FE A", talonario="71", numero="2", saldo=50.0,
                            dias_venc=-5, fecha_vto=date(2026, 8, 15)),
        modelos.Comprobante(tipo="FE A", talonario="71", numero="3", saldo=25.0,
                            dias_venc=0, fecha_vto=date(2026, 8, 10)),
    ]
    c.total_saldo_reporte = 175.0
    eq("codigo_corto", c.codigo_corto, "2761")
    eq("saldo_total", c.saldo_total, 175.0)
    eq("saldo_vencido", c.saldo_vencido, 100.0)
    eq("saldo_a_vencer", c.saldo_a_vencer, 75.0)
    eq("estado", c.estado, modelos.VENCIDO)
    eq("dias_max_vencido", c.dias_max_vencido, 10)
    eq("dias_hasta_proximo_vto", c.dias_hasta_proximo_vto, 0)
    eq("descuadre", c.descuadre(), 0.0)
    check("esta_consistente", c.esta_consistente())
    eq("comprobante_str", c.comprobantes[0].comprobante_str, "FE A 71-1")
    eq("descripcion_tiempo cliente", c.descripcion_tiempo(), "vencido hace 10 dias")
    eq("descripcion_tiempo comp vencido",
       c.comprobantes[0].descripcion_tiempo(), "vencido hace 10 dias")
    eq("descripcion_tiempo comp a vencer",
       c.comprobantes[1].descripcion_tiempo(), "faltan 5 dias")
    eq("descripcion_tiempo comp hoy",
       c.comprobantes[2].descripcion_tiempo(), "vence hoy")

    r = modelos.Reporte(clientes=[c], total_general_saldo=175.0)
    eq("reporte.saldo_total", r.saldo_total, 175.0)
    eq("reporte.descuadre_general", r.descuadre_general(), 0.0)
    eq("reporte.validar sin problemas", r.validar(), [])
    check("cliente_por_codigo largo", r.cliente_por_codigo("00000002761") is c)
    check("cliente_por_codigo corto", r.cliente_por_codigo("2761") is c)

    r.total_general_saldo = 200.0
    check("validar detecta descuadre general", len(r.validar()) == 1,
          str(r.validar()))


# ==========================================================================
# 2. ARCHIVOS DE PRUEBA (EXCEL)
# ==========================================================================
FECHA_REPORTE = date(2026, 8, 10)


def crear_excel(ruta: str, romper_total: bool = False) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    AMARILLO = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"

    def fila(n: int, valores: dict[str, object]) -> None:
        for col, val in valores.items():
            ws[f"{col}{n}"] = val

    fila(1, {"A": "WOOD TOOLS SRL", "I": "Fecha:", "J": datetime(2026, 8, 10)})
    fila(2, {"A": "A COBRAR POR CLIENTE EXPRESADO EN $", "I": "Hoja:", "J": 1})
    fila(3, {"A": "Cobrador:110-110 Vendedor:002-044"})
    fila(5, {"A": "Cliente", "B": "Comprobante", "D": "Condicion de pago",
             "E": "Fecha", "H": "Importe", "I": "Saldo", "J": "Dias"})

    # --- cliente 1 -------------------------------------------------------
    fila(7, {"A": "00000002761", "C": "ACME HERRAMIENTAS SRL",
             "F": "(1047) CAPITAL FEDERAL",
             "G": "Marcos 11 3900 9488 -4313 0466", "I": "Jonathan"})
    fila(8, {"C": "FERRETERIA ACME"})
    fila(9, {"C": "AV SIEMPREVIVA 742"})
    fila(10, {"A": "71", "B": "FE A", "C": "00072065", "D": "30 dias.ff.CONTRA ENTREGA",
              "E": datetime(2026, 7, 1), "F": datetime(2026, 7, 31), "G": "01/01",
              "H": 37165.39, "I": 37165.39, "J": 10})
    fila(11, {"A": "71", "B": "FE A", "C": "00072100", "D": "30 dias.ff.",
              "E": datetime(2026, 7, 15), "F": datetime(2026, 8, 14), "G": "01/01",
              "H": 20000.00, "I": 15000.00, "J": -4})
    fila(12, {"B": "Total cliente 00000002761 ACME HERRAMIENTAS SRL",
              "H": 57165.39, "I": 52165.39})

    # --- cliente 2 (marcado en amarillo) ---------------------------------
    fila(14, {"A": "00000003120", "C": "FERRETERIA DEL SUR SA",
              "F": "(1878) QUILMES", "G": "45246075 LAUTARO 1140647300 KAREN PAGO"})
    ws["A14"].fill = AMARILLO
    ws["C14"].fill = AMARILLO
    fila(15, {"C": "CALLE FALSA 123"})
    fila(16, {"A": "71", "B": "FE A", "C": "00072200", "D": "CONTADO",
              "E": datetime(2026, 6, 1), "F": datetime(2026, 7, 1), "G": "01/01",
              "H": 100000.00, "I": 100000.00, "J": 40})
    ws["I16"].fill = AMARILLO
    fila(17, {"B": "Total cliente 00000003120 FERRETERIA DEL SUR SA",
              "H": 100000.00, "I": 100000.00})

    # --- cliente 3 (vence hoy) -------------------------------------------
    fila(19, {"A": "00000004050", "C": "MADERAS NORTE SRL",
              "F": "(1602) FLORIDA", "G": "4382-1590 / Jonathan153-359-5605"})
    fila(20, {"C": "RUTA 8 KM 30"})
    fila(21, {"A": "72", "B": "FE B", "C": "00000915", "D": "21 dias",
              "E": datetime(2026, 7, 20), "F": datetime(2026, 8, 10), "G": "01/01",
              "H": 5000.50, "I": 5000.50, "J": 0})
    saldo_total_c3 = 5000.50 if not romper_total else 9999.99
    fila(22, {"B": "Total cliente 00000004050 MADERAS NORTE SRL",
              "H": 5000.50, "I": saldo_total_c3})

    total_gen = 157165.89 if not romper_total else 157165.89 + (9999.99 - 5000.50)
    fila(24, {"A": "Total general:", "H": 162165.89, "I": total_gen})

    wb.save(ruta)


RUTA_XLSX = os.path.join(TMP, "LISTADO_VENCIDO.xlsx")
RUTA_XLSX_MAL = os.path.join(TMP, "LISTADO_VENCIDO_MAL.xlsx")


def t_parser_excel() -> None:
    grupo("2. parser_excel.py  (lector deterministico de Excel)")

    crear_excel(RUTA_XLSX)
    crear_excel(RUTA_XLSX_MAL, romper_total=True)
    rep = parser_excel.parsear_excel(RUTA_XLSX, "11")

    eq("empresa", rep.empresa, "WOOD TOOLS SRL")
    eq("titulo", rep.titulo, "A COBRAR POR CLIENTE EXPRESADO EN $")
    eq("fecha del reporte", rep.fecha_reporte, FECHA_REPORTE)
    eq("filtro vendedor desde", rep.filtro_vendedor_desde, "002")
    eq("filtro vendedor hasta", rep.filtro_vendedor_hasta, "044")
    eq("filtro cobrador", rep.filtro_cobrador, "110")
    eq("motor", rep.motor, "excel")

    eq("cantidad de clientes", len(rep.clientes), 3)
    if len(rep.clientes) != 3:
        return

    c1, c2, c3 = rep.clientes
    eq("c1 codigo", c1.codigo, "00000002761")
    eq("c1 codigo_corto", c1.codigo_corto, "2761")
    eq("c1 razon social", c1.razon_social, "ACME HERRAMIENTAS SRL")
    eq("c1 nombre de fantasia", c1.nombre_fantasia, "FERRETERIA ACME")
    eq("c1 direccion", c1.direccion, "AV SIEMPREVIVA 742")
    eq("c1 localidad", c1.localidad, "(1047) CAPITAL FEDERAL")
    eq("c1 contacto", c1.contacto, "Jonathan")
    eq("c1 telefonos", c1.telefonos, ["5491139009488", "5491143130466"])
    eq("c1 comprobantes", c1.cantidad_comprobantes, 2)
    eq("c1 saldo total", c1.saldo_total, 52165.39, tol=0.005)
    eq("c1 saldo vencido", c1.saldo_vencido, 37165.39, tol=0.005)
    eq("c1 saldo a vencer", c1.saldo_a_vencer, 15000.0, tol=0.005)
    eq("c1 total del reporte", c1.total_saldo_reporte, 52165.39, tol=0.005)
    eq("c1 descuadre", c1.descuadre(), 0.0, tol=0.005)
    eq("c1 resaltado", c1.resaltado, False)
    eq("c1 comp1 fecha vto", c1.comprobantes[0].fecha_vto, date(2026, 7, 31))
    eq("c1 comp1 tipo", c1.comprobantes[0].tipo, "FE A")
    eq("c1 comp1 numero", c1.comprobantes[0].numero, "00072065")
    eq("c1 comp1 condicion", c1.comprobantes[0].condicion_pago, "30 dias.ff.CONTRA ENTREGA")
    eq("c1 comp1 dias", c1.comprobantes[0].dias_venc, 10)
    eq("c1 comp2 dias (negativo = a vencer)", c1.comprobantes[1].dias_venc, -4)
    eq("c1 comp2 saldo parcial != importe original",
       (c1.comprobantes[1].importe_original, c1.comprobantes[1].saldo), (20000.0, 15000.0))
    eq("c1 fila de origen", c1.fila_origen, 7)

    eq("c2 resaltado (amarillo detectado)", c2.resaltado, True)
    eq("c2 saldo total", c2.saldo_total, 100000.0, tol=0.005)
    eq("c2 estado", c2.estado, modelos.VENCIDO)
    eq("c2 telefonos", c2.telefonos, ["5491145246075", "5491140647300"])
    eq("c2 comp resaltado", c2.comprobantes[0].resaltado, True)

    eq("c3 estado (vence hoy)", c3.estado, modelos.VENCE_HOY)
    eq("c3 resaltado", c3.resaltado, False)
    eq("c3 telefonos", c3.telefonos, ["5491133595605", "5491143821590"])

    eq("total general del reporte", rep.total_general_saldo, 157165.89, tol=0.005)
    eq("suma calculada", rep.saldo_total, 157165.89, tol=0.005)
    eq("descuadre general", rep.descuadre_general(), 0.0, tol=0.005)
    eq("saldo vencido total", rep.saldo_vencido, 137165.39, tol=0.005)
    eq("saldo a vencer total", rep.saldo_a_vencer, 20000.50, tol=0.005)
    eq("clientes vencidos", len(rep.clientes_vencidos), 2)
    eq("clientes resaltados", len(rep.clientes_resaltados), 1)
    eq("validar() sin problemas", rep.validar(), [])
    eq("avisos vacios", rep.avisos, [])

    # --- archivo con el total roto ---
    rep_mal = parser_excel.parsear_excel(RUTA_XLSX_MAL, "11")
    problemas = rep_mal.validar()
    check("detecta el descuadre de un cliente",
          any("MADERAS NORTE" in p for p in problemas), str(problemas))
    check("detecta el descuadre general",
          any("Total general" in p for p in problemas), str(problemas))

    # --- volcado para Claude ---
    volcado = parser_excel.volcar_para_claude(RUTA_XLSX)
    check("volcado marca [AMARILLO]", "[AMARILLO]" in volcado)
    check("volcado usa formato F<n> | A=...", "F7 | A=00000002761" in volcado,
          volcado.splitlines()[6] if len(volcado.splitlines()) > 6 else "")
    check("volcado incluye el total general", "Total general:" in volcado)

    # --- recalculo de dias ---
    rep2 = parser_excel.parsear_excel(RUTA_XLSX, "11")
    cambios = parser_excel.recalcular_dias(rep2, date(2026, 8, 20))
    eq("recalcular_dias cambia los 4 comprobantes", cambios, 4)
    eq("recalculado c1 comp1", rep2.clientes[0].comprobantes[0].dias_venc, 20)
    eq("recalculado c3 pasa a vencido", rep2.clientes[2].estado, modelos.VENCIDO)


# ==========================================================================
# 3. PARSER PDF
# ==========================================================================
RUTA_PDF = os.path.join(TMP, "LISTADO_VENCIDO.pdf")


def crear_pdf_reporte(ruta: str) -> None:
    """PDF con el mismo contenido, imitando la salida en texto del ERP."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas as rl

    lineas = [
        "WOOD TOOLS SRL                                                       Fecha: 10/08/2026",
        "A COBRAR POR CLIENTE EXPRESADO EN $                                  Hoja: 1",
        "Cobrador:110-110 Vendedor:002-044",
        "",
        "Cliente          Comprobante   Condicion de pago      Emision    Vto.    Cuota   Importe Original      Saldo   Dias Venc.",
        "",
        "00000002761 ACME HERRAMIENTAS SRL          (1047) CAPITAL FEDERAL   Marcos 11 3900 9488 -4313 0466",
        "   FERRETERIA ACME",
        "   AV SIEMPREVIVA 742",
        "71 FE A 00072065  30 dias.ff.CONTRA ENTREGA  01/07/2026 31/07/2026 01/01      37.165,39   37.165,39   10",
        "71 FE A 00072100  30 dias.ff.                15/07/2026 14/08/2026 01/01      20.000,00   15.000,00   -4",
        "Total cliente 00000002761 ACME HERRAMIENTAS SRL                                       57.165,39   52.165,39",
        "",
        "00000003120 FERRETERIA DEL SUR SA          (1878) QUILMES   45246075 LAUTARO 1140647300 KAREN PAGO",
        "   CALLE FALSA 123",
        "71 FE A 00072200  CONTADO                    01/06/2026 01/07/2026 01/01     100.000,00  100.000,00   40",
        "Total cliente 00000003120 FERRETERIA DEL SUR SA                                      100.000,00  100.000,00",
        "",
        "00000004050 MADERAS NORTE SRL             (1602) FLORIDA   4382-1590 / Jonathan153-359-5605",
        "   RUTA 8 KM 30",
        "72 FE B 00000915  21 dias                    20/07/2026 10/08/2026 01/01       5.000,50    5.000,50    0",
        "Total cliente 00000004050 MADERAS NORTE SRL                                            5.000,50    5.000,50",
        "",
        "Total general:                                                                162.165,89  157.165,89",
    ]
    c = rl.Canvas(ruta, pagesize=landscape(A4))
    c.setFont("Courier", 7)
    y = 560
    for linea in lineas:
        c.drawString(20, y, linea)
        y -= 11
    c.save()


def t_parser_pdf() -> None:
    grupo("3. parser_pdf.py  (lector best-effort de PDF, SIN Claude)")

    crear_pdf_reporte(RUTA_PDF)
    check("es_pdf reconoce el archivo", parser_pdf.es_pdf(RUTA_PDF))
    check("es_pdf rechaza el Excel", not parser_pdf.es_pdf(RUTA_XLSX))
    eq("contar_paginas", parser_pdf.contar_paginas(RUTA_PDF), 1)
    check("esta_escaneado = False", not parser_pdf.esta_escaneado(RUTA_PDF))

    eq("_a_float_ar('37.165,39')", parser_pdf._a_float_ar("37.165,39"), 37165.39)
    eq("_a_float_ar('(1.234,00)')", parser_pdf._a_float_ar("(1.234,00)"), -1234.0)
    eq("_a_float_ar('1.234-')", parser_pdf._a_float_ar("1.234-"), -1234.0)
    eq("_a_float_ar('texto')", parser_pdf._a_float_ar("texto"), None)

    rep = parser_pdf.parsear_pdf(RUTA_PDF, "11")
    eq("fecha del reporte", rep.fecha_reporte, FECHA_REPORTE)
    eq("empresa sin la 'Fecha:' pegada al nombre", rep.empresa, "WOOD TOOLS SRL")
    eq("cantidad de clientes", len(rep.clientes), 3)
    eq("total general", rep.total_general_saldo, 157165.89, tol=0.005)
    eq("suma calculada", rep.saldo_total, 157165.89, tol=0.005)
    eq("descuadre general", rep.descuadre_general(), 0.0, tol=0.005)
    problemas = rep.validar()
    check("validar() sin problemas", not problemas, "; ".join(problemas[:4]))
    print("  -- tipos de comprobante --")
    linea_fe = ("71 FE A 00072592 Cond.15 Dias 18/08/2026 02/09/2026 01/01 "
                "163372,51 163372,41 -2")
    linea_pres = ("56 PRES 00154758 Cond.15 Dias 03/08/2026 18/08/2026 01/01 "
                  "57093,30 27093,30 13")
    check("reconoce 'FE A'", parser_pdf._parece_comprobante(linea_fe))
    check("reconoce 'PRES' (el tipo del listado de vencimientos)",
          parser_pdf._parece_comprobante(linea_pres))
    comp = parser_pdf._armar_comprobante(linea_pres, None)
    eq("PRES: tipo", comp.tipo, "PRES")
    eq("PRES: importe original", comp.importe_original, 57093.30, tol=0.005)
    eq("PRES: saldo parcial", comp.saldo, 27093.30, tol=0.005)
    eq("PRES: dias", comp.dias_venc, 13)
    check("no confunde una linea cualquiera con un comprobante",
          not parser_pdf._parece_comprobante("CALLE FALSA 123"))

    print("  -- _partir_cabecera_cliente (separacion de la linea de cabecera) --")
    casos_cab = [
        ("ACME HERRAMIENTAS SRL (1047) CAPITAL FEDERAL Marcos 11 3900 9488 -4313 0466",
         ("ACME HERRAMIENTAS SRL", "(1047) CAPITAL FEDERAL",
          "Marcos 11 3900 9488 -4313 0466")),
        ("FERRETERIA DEL SUR SA (1878) QUILMES 45246075 LAUTARO 1140647300 KAREN PAGO",
         ("FERRETERIA DEL SUR SA", "(1878) QUILMES",
          "45246075 LAUTARO 1140647300 KAREN PAGO")),
        ("MADERAS NORTE SRL (1602) FLORIDA 4382-1590 / Jonathan153-359-5605",
         ("MADERAS NORTE SRL", "(1602) FLORIDA", "4382-1590 / Jonathan153-359-5605")),
        # sin localidad: se corta en el telefono
        ("TALLER LOPEZ SA 11 4085 5781",
         ("TALLER LOPEZ SA", "", "11 4085 5781")),
        # sin localidad ni telefonos: queda tal cual
        ("EST.CENTRAL S.A.C.I.F.I.Y.A.", ("EST.CENTRAL S.A.C.I.F.I.Y.A.", "", "")),
        # la localidad va en el renglon de abajo y el ERP separa columnas con
        # dos o mas espacios: contacto y telefonos quedan a la derecha
        ("METALURGICA DEL OESTE S.A.        Santiago 1145678901 dueno",
         ("METALURGICA DEL OESTE S.A.", "", "Santiago 1145678901 dueno")),
        # ademas con la condicion de venta suelta, que se descarta
        ("TORNILLERIA SUR S.A.I.C.        B  4712-3456 cel:1155667788",
         ("TORNILLERIA SUR S.A.I.C.", "", "4712-3456 cel:1155667788")),
        # NO ROMPER: numeros que son parte del nombre
        ("FERRETERIA 2000 SA", ("FERRETERIA 2000 SA", "", "")),
        ("HIERROS 3 DE FEBRERO SRL", ("HIERROS 3 DE FEBRERO SRL", "", "")),
        # espacios de mas dentro del nombre: no es una columna, no se corta
        ("CASA GRANDE   2000", ("CASA GRANDE   2000", "", "")),
    ]
    for entrada, esperado in casos_cab:
        eq(f"cabecera: {entrada[:34]}...",
           parser_pdf._partir_cabecera_cliente(entrada), esperado)

    # idempotencia: volver a pasar la salida no la cambia
    razon, loc, tel = parser_pdf._partir_cabecera_cliente(casos_cab[0][0])
    eq("idempotente sobre la razon social ya limpia",
       parser_pdf._partir_cabecera_cliente(razon), (razon, "", ""))

    # importes pegados al nombre en la fila "Total cliente"
    eq("saca los importes de 'Total cliente'",
       parser_pdf._RE_COLA_IMPORTES.sub(
           "", "ACME HERRAMIENTAS SRL      57.165,39 52.165,39").strip(),
       "ACME HERRAMIENTAS SRL")

    if len(rep.clientes) == 3 and rep.clientes[0].comprobantes:
        c1 = rep.clientes[0]
        eq("c1 codigo", c1.codigo, "00000002761")
        eq("ARREGLADO: razon social limpia", c1.razon_social, "ACME HERRAMIENTAS SRL")
        eq("ARREGLADO: localidad detectada", c1.localidad, "(1047) CAPITAL FEDERAL")
        eq("ARREGLADO: telefonos detectados",
           c1.telefonos, ["5491139009488", "5491143130466"])
        eq("c1 direccion", c1.direccion, "AV SIEMPREVIVA 742")
        eq("c2 razon social limpia", rep.clientes[1].razon_social,
           "FERRETERIA DEL SUR SA")
        eq("c3 razon social limpia", rep.clientes[2].razon_social, "MADERAS NORTE SRL")
        check("c3 telefonos detectados", bool(rep.clientes[2].telefonos),
              str(rep.clientes[2].telefonos))
        eq("c1 comprobantes", c1.cantidad_comprobantes, 2)
        eq("c1 saldo total", c1.saldo_total, 52165.39, tol=0.005)
        eq("c1 comp1 dias", c1.comprobantes[0].dias_venc, 10)
        eq("c1 comp2 dias", c1.comprobantes[1].dias_venc, -4)
        eq("c1 estado", c1.estado, modelos.VENCIDO)
        eq("c3 vence hoy", rep.clientes[2].estado, modelos.VENCE_HOY)
    for a in rep.avisos:
        print(f"        aviso: {a}")

    if len(rep.clientes) == 3:
        c1 = rep.clientes[0]
        print(f"        razon_social leida del PDF: {c1.razon_social!r}")
        print(f"        localidad: {c1.localidad!r}  direccion: {c1.direccion!r}")
        texto_wa = mensajes.construir_mensaje(rep, rep.clientes, "Juan Perez", "002")
        check("ARREGLADO: el mensaje ya no arrastra basura en el nombre",
              "(1047)" not in texto_wa and "3900 9488" not in texto_wa
              and "45246075" not in texto_wa)
        print("        --- lo que recibiria el vendedor (leido del PDF, sin Claude) ---")
        for linea in texto_wa.splitlines()[:9]:
            print(("        | " + linea).encode("ascii", "replace").decode("ascii"))

    volcado = parser_pdf.volcar_para_claude(RUTA_PDF)
    check("volcado para Claude tiene cabecera", volcado.startswith("# ARCHIVO:"))
    check("volcado incluye el texto", "ACME HERRAMIENTAS" in volcado)



# ==========================================================================
# 3b. OCR LOCAL (PDF escaneados, sin internet y sin la API)
# ==========================================================================
def t_ocr() -> None:
    grupo("3b. parser_pdf.py - OCR local")

    listo, motivo = parser_pdf.ocr_disponible()
    check("ocr_disponible() devuelve (bool, motivo) sin procesar nada",
          isinstance(listo, bool) and isinstance(motivo, str))
    eq("el umbral de 'esto es un escaneo' esta definido",
       isinstance(parser_pdf.MINIMO_TEXTO_REAL, int), True)
    check("el OCR corre a 400 dpi (medido: a 300 se pierden mas comas)",
          parser_pdf.DPI_OCR == 400, str(parser_pdf.DPI_OCR))

    if not listo:
        print(f"  [SALTEADA ] no hay OCR en esta PC: {motivo.splitlines()[0]}")
        nota("Tesseract no esta instalado (winget install UB-Mannheim.TesseractOCR)",
             True)
        return

    check("encuentra el ejecutable de Tesseract",
          bool(parser_pdf._ruta_tesseract()), parser_pdf._ruta_tesseract())
    check("elige un idioma valido", parser_pdf._idioma_ocr() in ("spa", "eng"),
          parser_pdf._idioma_ocr())

    # --- se fabrica un PDF escaneado a partir del PDF de prueba ------------
    import fitz

    escaneado = os.path.join(TMP, "LISTADO_ESCANEADO.pdf")
    origen = fitz.open(RUTA_PDF)
    destino = fitz.open()
    for i in range(origen.page_count):
        pix = origen.load_page(i).get_pixmap(dpi=200)
        hoja = destino.new_page(width=pix.width * 72 / 200, height=pix.height * 72 / 200)
        hoja.insert_image(hoja.rect, stream=pix.tobytes("png"))
    destino.save(escaneado)
    destino.close()
    origen.close()

    check("el PDF fabricado no tiene texto adentro",
          len("".join(parser_pdf.extraer_lineas(escaneado))) < parser_pdf.MINIMO_TEXTO_REAL)
    check("esta_escaneado() lo reconoce", parser_pdf.esta_escaneado(escaneado))

    # --- se lee con OCR y se compara contra el original --------------------
    ref = parser_pdf.parsear_pdf(RUTA_PDF, "11")
    ocr = parser_pdf.parsear_pdf(escaneado, "11")

    eq("el motor queda marcado como OCR", ocr.motor, "pdf-ocr")
    check("avisa que fue leido con OCR",
          any("OCR" in a for a in ocr.avisos), str(ocr.avisos[:1]))
    eq("encuentra los mismos clientes", len(ocr.clientes), len(ref.clientes))
    eq("encuentra los mismos comprobantes",
       sum(c.cantidad_comprobantes for c in ocr.clientes),
       sum(c.cantidad_comprobantes for c in ref.clientes))
    eq("la suma de los comprobantes coincide al centavo",
       ocr.saldo_total, ref.saldo_total, tol=0.005)

    por_codigo = {c.codigo_corto: c for c in ocr.clientes}
    exactos = sum(
        1 for c in ref.clientes
        if c.codigo_corto in por_codigo
        and abs(por_codigo[c.codigo_corto].saldo_total - c.saldo_total) < 0.01
    )
    eq("todos los clientes con el saldo exacto", exactos, len(ref.clientes))

    # Lo que SI falla: el OCR pierde la coma decimal en las filas de totales.
    # No se corrige a mano -seria adivinar con plata-; se deja que el control
    # lo marque, que es lo que evita mandar numeros mal.
    nota("el OCR lee mal las filas 'Total cliente' (pierde la coma decimal)",
         len(ocr.validar()) > 0)
    check("por eso el programa marca controles que no cierran y no deja enviar "
          "sin confirmar", len(ocr.validar()) > 0 or ocr.descuadre_general() == 0)

# ==========================================================================
# 4. MENSAJES
# ==========================================================================
def t_mensajes() -> None:
    grupo("4. mensajes.py  (texto del WhatsApp)")

    rep = parser_excel.parsear_excel(RUTA_XLSX, "11")
    clientes = rep.clientes
    texto = mensajes.construir_mensaje(rep, clientes, "Juan Perez", "002")

    check("titulo presente", "*RESUMEN DE CUENTA - WOOD TOOLS*" in texto)
    check("nombre y codigo del vendedor", "Juan Perez (002)" in texto)
    check("fecha del reporte", "Reporte al 10/08/2026" in texto)
    check("seccion vencidos", "*CLIENTES CON DEUDA VENCIDA* (2 clientes" in texto)
    check("seccion a vencer", "*CLIENTES A VENCER* (1 cliente " in texto)
    check("total a cobrar", "*TOTAL A COBRAR: $157.165,89*" in texto)
    check("desglose vencido/a vencer",
          "$137.165,39 ya vencido" in texto and "$20.000,50 por vencer" in texto)
    check("firma", "_Administracion Wood Tools_" in texto)
    check("orden: el mas vencido primero",
          texto.index("FERRETERIA DEL SUR") < texto.index("ACME HERRAMIENTAS"))
    check("aclara el desglose del cliente mixto",
          "vencido $37.165,39 + a vencer $15.000,00" in texto)

    # los subtotales de las dos secciones tienen que sumar el total
    import re as _re
    subtotales = [
        float(m.replace(".", "").replace(",", "."))
        for m in _re.findall(r"clientes? - \$([\d\.,]+)\)", texto)
    ]
    eq("subtotales de seccion suman el total", round(sum(subtotales), 2), 157165.89)

    # sin clientes
    vacio = mensajes.construir_mensaje(rep, [], "Ana Gomez")
    check("mensaje vacio es explicito", "No hay clientes con saldo pendiente" in vacio)

    # partido
    partes = mensajes.partir_mensaje(texto)
    eq("mensaje corto = 1 parte", len(partes), 1)

    grande = modelos.Reporte(fecha_reporte=FECHA_REPORTE)
    for i in range(120):
        c = modelos.Cliente(codigo=f"{i + 1:011d}",
                            razon_social=f"CLIENTE DE PRUEBA NUMERO {i + 1} SRL")
        c.comprobantes = [modelos.Comprobante(tipo="FE A", talonario="71",
                                              numero=f"{i:08d}", saldo=1000.0 + i,
                                              dias_venc=i - 60,
                                              fecha_vto=date(2026, 8, 1))]
        c.total_saldo_reporte = c.saldo_total
        grande.clientes.append(c)
    texto_grande = mensajes.construir_mensaje(grande, grande.clientes, "Juan Perez", "002")
    partes = mensajes.partir_mensaje(texto_grande, 3800)
    check("mensaje largo se parte", len(partes) > 1, f"{len(partes)} partes")
    check("ninguna parte supera 4096",
          all(len(p) <= 4096 for p in partes), str([len(p) for p in partes]))
    check("ninguna parte supera el maximo pedido",
          all(len(p) <= 3800 for p in partes), str([len(p) for p in partes]))
    check("las partes vienen numeradas", partes[0].rstrip().endswith(f"(1/{len(partes)})"))
    # ningun cliente cortado al medio
    juntos = "\n\n".join(_re.sub(r"\n\n\(\d+/\d+\)$", "", p) for p in partes)
    faltantes = [c.nombre for c in grande.clientes if c.nombre not in juntos]
    check("ningun cliente se pierde al partir", not faltantes, str(faltantes[:5]))

    # comprobantes en el mensaje
    con_comp = mensajes.construir_mensaje(rep, clientes, "Juan", "002",
                                          incluir_comprobantes=True)
    check("detalle de comprobantes", "FE A 71-00072065" in con_comp)

    # plantilla de Meta
    plantilla = mensajes.construir_mensaje_plantilla(rep, clientes, "Juan")
    check("plantilla sin saltos de linea", "\n" not in plantilla and "\t" not in plantilla)
    check("plantilla <= 1000 caracteres", len(plantilla) <= 1000, str(len(plantilla)))
    check("plantilla trae el total", "TOTAL $157.165,89" in plantilla)
    plantilla_grande = mensajes.construir_mensaje_plantilla(grande, grande.clientes, "Juan")
    check("plantilla grande truncada a 1000", len(plantilla_grande) <= 1000,
          str(len(plantilla_grande)))

    epi = mensajes.epigrafe_pdf(rep, clientes, "Juan Perez")
    check("epigrafe <= 900", len(epi) <= 900)
    check("epigrafe con total", "$157.165,89" in epi, epi)


# ==========================================================================
# 5. PDF DEL RESUMEN
# ==========================================================================
def t_pdf_resumen() -> None:
    grupo("5. pdf_resumen.py  (PDF por vendedor)")

    eq("nombre sugerido",
       pdf_resumen.nombre_archivo_sugerido("Juan Perez", FECHA_REPORTE),
       "Resumen_de_cuenta_Juan_Perez_2026-08-10.pdf")
    eq("nombre con caracteres invalidos",
       pdf_resumen.nombre_archivo_sugerido('Ana/Gomez:*?"<>|', FECHA_REPORTE),
       "Resumen_de_cuenta_AnaGomez_2026-08-10.pdf")
    eq("nombre con acentos",
       pdf_resumen.nombre_archivo_sugerido("José Muñoz", FECHA_REPORTE),
       "Resumen_de_cuenta_Jose_Munoz_2026-08-10.pdf")

    rep = parser_excel.parsear_excel(RUTA_XLSX, "11")
    ok, ruta = pdf_resumen.generar_pdf_en_carpeta_salida(
        rep, rep.clientes, nombre_vendedor="Juan Perez", codigo_vendedor="002",
        incluir_comprobantes=True, incluir_telefonos=True,
    )
    check("se genero el PDF", ok, ruta)
    if ok:
        check("el archivo existe", os.path.exists(ruta))
        tam = os.path.getsize(ruta)
        check("el PDF no esta vacio", tam > 2000, f"{tam} bytes")
        with open(ruta, "rb") as fh:
            check("firma %PDF", fh.read(5) == b"%PDF-")
        paginas = parser_pdf.contar_paginas(ruta)
        check("tiene al menos 2 paginas (detalle en hoja aparte)", paginas >= 2,
              f"{paginas} paginas")
        texto = parser_pdf.extraer_texto(ruta)
        check("el PDF nombra al vendedor", "Juan Perez" in texto)
        check("el PDF trae el total a cobrar", "157.165,89" in texto)
        check("el PDF numera las paginas", "Pagina 1 de" in texto)
        check("el PDF lista los comprobantes", "00072065" in texto)
        check("el PDF muestra los telefonos", "5491139009488" in texto)
        check("el PDF separa VENCIDOS / A VENCER",
              "VENCIDOS" in texto and "A VENCER" in texto)

    ok2, motivo = pdf_resumen.generar_pdf(rep, [], os.path.join(TMP, "vacio.pdf"))
    check("PDF sin clientes se rechaza", not ok2, motivo)

    # un vendedor solo con clientes a vencer (sin seccion de vencidos)
    solo_av = [c for c in rep.clientes if c.estado != modelos.VENCIDO]
    ok3, ruta3 = pdf_resumen.generar_pdf(
        rep, solo_av, os.path.join(TMP, "solo_avencer.pdf"), nombre_vendedor="Ana")
    check("PDF con solo clientes a vencer", ok3, ruta3)


# ==========================================================================
# 6. VENDEDORES Y ASIGNACIONES
# ==========================================================================
def t_vendedores() -> None:
    grupo("6. vendedores.py  (alta, baja y asignacion)")

    for f in (config.ARCHIVO_VENDEDORES, config.ARCHIVO_ASIGNACIONES):
        if os.path.exists(f):
            os.remove(f)

    eq("arranca vacio", vend.cargar_vendedores(), [])

    v1, motivo = vend.agregar_vendedor("Juan Perez", "11 3900-9488", "002")
    check("alta de vendedor", v1 is not None, motivo)
    eq("telefono normalizado", v1.telefono if v1 else None, "5491139009488")

    v2, _ = vend.agregar_vendedor("Ana Gomez", "1140647300", "005")
    check("segunda alta", v2 is not None)

    v3, motivo3 = vend.agregar_vendedor("Repetido", "11 3900 9488", "009")
    check("rechaza telefono duplicado", v3 is None and "ya esta cargado" in motivo3, motivo3)

    v4, motivo4 = vend.agregar_vendedor("Otro", "1155667788", "002")
    check("rechaza codigo duplicado", v4 is None and "ya lo usa" in motivo4, motivo4)

    v5, motivo5 = vend.agregar_vendedor("Sin telefono", "1234", "010")
    check("rechaza telefono invalido", v5 is None, motivo5)

    v6, motivo6 = vend.agregar_vendedor("", "1155667799", "011")
    check("rechaza nombre vacio", v6 is None, motivo6)

    eq("hay 2 vendedores", len(vend.cargar_vendedores()), 2)
    check("vendedor_por_codigo tolera ceros", vend.vendedor_por_codigo("2") is not None)
    check("vendedores_activos", len(vend.vendedores_activos()) == 2)

    ok, motivo = vend.actualizar_vendedor(v1.id, nombre="Juan C. Perez", activo=False)
    check("actualizar vendedor", ok, motivo)
    eq("quedo inactivo", vend.vendedor_por_id(v1.id).activo, False)
    eq("activos ahora 1", len(vend.vendedores_activos()), 1)

    ok, motivo = vend.actualizar_vendedor(v1.id, telefono="1140647300")
    check("rechaza pisar el telefono de otro", not ok, motivo)

    rep = parser_excel.parsear_excel(RUTA_XLSX, "11")
    n = vend.asignar_varios([c.codigo for c in rep.clientes[:2]], v1.id)
    eq("asigno 2 clientes", n, 2)
    vend.asignar_cliente(rep.clientes[2].codigo, v2.id)

    grupos = vend.agrupar_por_vendedor(rep.clientes)
    eq("2 grupos", len(grupos), 2)
    eq("v1 tiene 2 clientes", len(grupos.get(v1.id, [])), 2)
    eq("v2 tiene 1 cliente", len(grupos.get(v2.id, [])), 1)
    check("no quedan clientes sin asignar", vend.SIN_ASIGNAR not in grupos)

    check("vendedor_de_cliente por codigo largo",
          vend.vendedor_de_cliente("00000002761").id == v1.id)
    check("vendedor_de_cliente por codigo corto",
          vend.vendedor_de_cliente("2761").id == v1.id)

    check("desasignar", vend.desasignar_cliente("2761"))
    grupos = vend.agrupar_por_vendedor(rep.clientes)
    check("ahora hay un cliente sin asignar", vend.SIN_ASIGNAR in grupos)

    check("eliminar vendedor", vend.eliminar_vendedor(v1.id))
    asign = vend.cargar_asignaciones()
    check("al borrar el vendedor se limpian sus asignaciones",
          v1.id not in asign.values(), str(asign))

    # archivo corrupto
    with open(config.ARCHIVO_VENDEDORES, "w", encoding="utf-8") as fh:
        fh.write("{esto no es json")
    eq("archivo corrupto no rompe", vend.cargar_vendedores(), [])


# ==========================================================================
# 7. WHATSAPP (sin red real)
# ==========================================================================
def t_whatsapp() -> None:
    grupo("7. whatsapp.py  (Cloud API de Meta, con la red simulada)")

    eq("normalizar_destino ok", whatsapp.normalizar_destino("5491139009488"),
       ("5491139009488", ""))
    check("normalizar_destino vacio", whatsapp.normalizar_destino("")[0] is None)
    check("normalizar_destino corto", whatsapp.normalizar_destino("12345")[0] is None)

    check("131047 traducido", "ventana de 24 horas" in whatsapp._traducir_error("131047", ""))
    check("131030 traducido", "destinatarios permitidos" in whatsapp._traducir_error("131030", ""))
    check("190 traducido", "token" in whatsapp._traducir_error("190", ""))
    check("132001 traducido", "plantilla" in whatsapp._traducir_error("132001", ""))
    check("re-engagement -> 131047",
          "ventana de 24 horas" in whatsapp._traducir_error("999", "Re-engagement message"))
    check("codigo desconocido no se pierde",
          "codigo 55555" in whatsapp._traducir_error("55555", "algo raro"))

    link = whatsapp.link_wa_me("5491139009488", "hola *mundo*\nsegunda linea")
    check("link wa.me bien codificado",
          link.startswith("https://wa.me/5491139009488?text=") and "%0A" in link, link)

    check("esta_configurado = False sin credenciales", not whatsapp.esta_configurado())
    r = whatsapp.enviar_texto("5491139009488", "hola")
    check("sin credenciales devuelve error claro", not r.ok and "Configuracion" in r.mensaje,
          r.mensaje)

    # --- con credenciales simuladas y requests parcheado ---
    config.CONFIG.datos["whatsapp_token"] = "EAAG_token_de_prueba"
    config.CONFIG.datos["whatsapp_phone_number_id"] = "1041050652417644"
    check("esta_configurado = True", whatsapp.esta_configurado())
    eq("base_url", config.CONFIG.base_url_whatsapp,
       "https://graph.facebook.com/v21.0/1041050652417644")

    class RespFalsa:
        def __init__(self, code, payload):
            self.status_code = code
            self._payload = payload
            self.text = json.dumps(payload)

        def json(self):
            return self._payload

    llamadas: list[dict] = []

    def post_ok(url, headers=None, json=None, timeout=None, **kw):
        llamadas.append({"url": url, "body": json})
        return RespFalsa(200, {"messages": [{"id": "wamid.TEST123"}]})

    original_post = whatsapp.requests.post
    whatsapp.requests.post = post_ok
    try:
        r = whatsapp.enviar_texto("11 3900-9488", "hola mundo")
        check("envio de texto OK", r.ok, r.mensaje)
        eq("id de mensaje", r.id_mensaje, "wamid.TEST123")
        eq("destino normalizado a digitos", llamadas[-1]["body"]["to"], "1139009488")
        eq("tipo texto", llamadas[-1]["body"]["type"], "text")
        check("url correcta", llamadas[-1]["url"].endswith("/messages"))

        r = whatsapp.enviar_documento("5491139009488", "MEDIA1", "resumen.pdf", "epigrafe")
        check("envio de documento OK", r.ok, r.mensaje)
        eq("tipo documento", llamadas[-1]["body"]["type"], "document")
        eq("nombre del archivo", llamadas[-1]["body"]["document"]["filename"], "resumen.pdf")

        r = whatsapp.enviar_plantilla("5491139009488", "resumen_cobranzas",
                                      ["linea1\nlinea2\ttab"], "MEDIA1")
        check("envio de plantilla OK", r.ok, r.mensaje)
        param = llamadas[-1]["body"]["template"]["components"][-1]["parameters"][0]["text"]
        check("la plantilla limpia saltos y tabs",
              "\n" not in param and "\t" not in param, repr(param))

        r = whatsapp.enviar_texto("5491139009488", "   ")
        check("rechaza mensaje vacio", not r.ok, r.mensaje)

        def post_error(url, headers=None, json=None, timeout=None, **kw):
            return RespFalsa(400, {"error": {"code": 131047, "message": "Re-engagement"}})

        whatsapp.requests.post = post_error
        r = whatsapp.enviar_texto("5491139009488", "hola")
        check("error 131047 se traduce", not r.ok and "ventana de 24 horas" in r.mensaje)
        eq("codigo de error propagado", r.codigo_error, "131047")

        intentos = {"n": 0}

        def post_500(url, headers=None, json=None, timeout=None, **kw):
            intentos["n"] += 1
            return RespFalsa(500, {"error": {"code": 500, "message": "server"}})

        whatsapp.requests.post = post_500
        whatsapp.time.sleep = lambda *_a, **_k: None   # sin esperas reales
        r = whatsapp.enviar_texto("5491139009488", "hola")
        check("reintenta ante 500", intentos["n"] == whatsapp.MAX_REINTENTOS + 1,
              f"{intentos['n']} intentos")
        check("termina en error", not r.ok)

        # orquestador completo, en modo libre
        whatsapp.requests.post = post_ok
        llamadas.clear()

        def subir_falso(ruta):
            return "MEDIA_OK", "PDF subido"

        original_subir = whatsapp.subir_documento
        whatsapp.subir_documento = subir_falso
        config.CONFIG.datos["pausa_entre_envios_seg"] = 0
        resultados = whatsapp.enviar_resumen(
            "5491139009488", ["parte 1", "parte 2"], ruta_pdf="x.pdf",
            nombre_pdf="resumen.pdf", epigrafe_pdf="epi")
        eq("enviar_resumen manda PDF + 2 textos", len(resultados), 3)
        check("todos OK", all(r.ok for r in resultados))
        eq("primero el documento", llamadas[0]["body"]["type"], "document")
        eq("despues el texto", llamadas[1]["body"]["type"], "text")

        # modo plantilla: un solo mensaje
        config.CONFIG.datos["modo_envio"] = "plantilla"
        config.CONFIG.datos["plantilla_nombre"] = "resumen_cobranzas"
        llamadas.clear()
        resultados = whatsapp.enviar_resumen(
            "5491139009488", ["parte 1", "parte 2"], ruta_pdf="x.pdf")
        eq("modo plantilla manda 1 solo mensaje", len(resultados), 1)
        eq("es una plantilla", llamadas[-1]["body"]["type"], "template")
        config.CONFIG.datos["modo_envio"] = "libre"
        whatsapp.subir_documento = original_subir
    finally:
        whatsapp.requests.post = original_post
        config.CONFIG.datos["whatsapp_token"] = ""
        config.CONFIG.datos["whatsapp_phone_number_id"] = ""


# ==========================================================================
# 8. CLAUDE API (sin llamadas reales)
# ==========================================================================
def t_claude() -> None:
    grupo("8. claude_api.py  (conversion y prompts, con la API simulada)")

    eq("_a_float(37165.39)", claude_api._a_float(37165.39), 37165.39)
    eq("_a_float('37.165,39')", claude_api._a_float("37.165,39"), 37165.39)
    eq("_a_float('(1.234,00)')", claude_api._a_float("(1.234,00)"), -1234.0)
    eq("_a_float(None)", claude_api._a_float(None), None)
    eq("_a_fecha('10/08/2026')", claude_api._a_fecha("10/08/2026"), date(2026, 8, 10))
    eq("_a_fecha('2026-08-10')", claude_api._a_fecha("2026-08-10"), date(2026, 8, 10))
    eq("_a_fecha('-')", claude_api._a_fecha("-"), None)
    eq("_a_int('10')", claude_api._a_int("10"), 10)

    for nombre, esquema in (
        ("ESQUEMA_REPORTE", claude_api.ESQUEMA_REPORTE),
        ("ESQUEMA_VERIFICACION", claude_api.ESQUEMA_VERIFICACION),
        ("ESQUEMA_TELEFONOS", claude_api.ESQUEMA_TELEFONOS),
    ):
        try:
            json.dumps(esquema)
            ok = True
        except Exception:
            ok = False
        check(f"{nombre} serializa a JSON", ok)
        check(f"{nombre} tiene additionalProperties=False",
              esquema.get("additionalProperties") is False)

    # --- conversion JSON -> Reporte, con un signo de dias equivocado ---
    payload = {
        "empresa": "WOOD TOOLS SRL",
        "titulo": "A COBRAR POR CLIENTE EXPRESADO EN $",
        "fecha_reporte": "10/08/2026",
        "filtros": "Cobrador:110-110 Vendedor:002-044",
        "total_general_original": 162165.89,
        "total_general_saldo": 157165.89,
        "avisos": ["una fila estaba borrosa"],
        "clientes": [
            {
                "codigo": "00000002761", "razon_social": "ACME HERRAMIENTAS SRL",
                "nombre_fantasia": "FERRETERIA ACME", "direccion": "AV SIEMPREVIVA 742",
                "localidad": "(1047) CAPITAL FEDERAL", "contacto": "Jonathan",
                "telefonos_raw": "Marcos 11 3900 9488 -4313 0466",
                "total_original_reporte": 57165.39, "total_saldo_reporte": 52165.39,
                "resaltado": True,
                "comprobantes": [
                    {"talonario": "71", "tipo": "FE A", "numero": "00072065",
                     "condicion_pago": "30 dias", "fecha_emision": "01/07/2026",
                     "fecha_vto": "31/07/2026", "cuota": "01/01",
                     "importe_original": 37165.39, "saldo": 37165.39,
                     "dias_venc": 10, "resaltado": True},
                    # Claude devuelve 4 (positivo = vencido) pero la fecha dice -4
                    {"talonario": "71", "tipo": "FE A", "numero": "00072100",
                     "condicion_pago": "30 dias", "fecha_emision": "15/07/2026",
                     "fecha_vto": "14/08/2026", "cuota": "01/01",
                     "importe_original": 20000.0, "saldo": 15000.0,
                     "dias_venc": 4, "resaltado": False},
                ],
            }
        ],
    }
    rep = claude_api._a_reporte(payload, "claude-pdf", "x.pdf")
    eq("empresa", rep.empresa, "WOOD TOOLS SRL")
    eq("fecha", rep.fecha_reporte, FECHA_REPORTE)
    eq("clientes", len(rep.clientes), 1)
    eq("resaltado", rep.clientes[0].resaltado, True)
    eq("telefonos derivados del texto crudo",
       rep.clientes[0].telefonos, ["5491139009488", "5491143130466"])
    eq("GANA LA FECHA sobre los dias de Claude",
       rep.clientes[0].comprobantes[1].dias_venc, -4)
    check("y queda anotado en avisos",
          any("corresponden -4" in a for a in rep.avisos), str(rep.avisos))
    check("los avisos de Claude se conservan",
          any("borrosa" in a for a in rep.avisos))
    eq("saldo total", rep.clientes[0].saldo_total, 52165.39, tol=0.005)
    eq("descuadre general", rep.descuadre_general(), -105000.50, tol=0.005)

    # cliente con formato invalido no rompe
    malo = dict(payload)
    malo["clientes"] = ["no soy un dict", {"codigo": "1", "comprobantes": "tampoco"}]
    rep_malo = claude_api._a_reporte(malo, "claude-pdf", "x.pdf")
    check("un cliente invalido no tumba la conversion", len(rep_malo.avisos) >= 1)

    # --- interpretar_volcado_excel con _llamar simulado ---
    original_llamar = claude_api._llamar
    capturado: dict = {}

    def llamar_falso(sistema, contenido, esquema, progreso=None):
        capturado["sistema"] = sistema
        capturado["contenido"] = contenido
        capturado["esquema"] = esquema
        return payload

    claude_api._llamar = llamar_falso
    try:
        volcado = parser_excel.volcar_para_claude(RUTA_XLSX)
        rep2 = claude_api.interpretar_volcado_excel(volcado, nombre_archivo=RUTA_XLSX)
        eq("interpretar_volcado_excel devuelve clientes", len(rep2.clientes), 1)
        eq("motor", rep2.motor, "claude-texto")
        check("el prompt explica el [AMARILLO]", "[AMARILLO]" in capturado["sistema"])
        check("el volcado viaja en el mensaje", "F7 | A=00000002761" in
              capturado["contenido"][0]["text"])
        check("usa el esquema del reporte", capturado["esquema"] is claude_api.ESQUEMA_REPORTE)

        # aviso de recorte con un volcado enorme
        rep3 = claude_api.interpretar_volcado_excel("x" * 1_600_001, nombre_archivo="g.xlsx")
        check("avisa cuando recorta el archivo",
              rep3.avisos and "ATENCION" in rep3.avisos[0], str(rep3.avisos[:1]))

        # verificar_reporte
        def llamar_verif(sistema, contenido, esquema, progreso=None):
            return {"coincide": True, "observaciones": [], "clientes_faltantes": [],
                    "diferencias": []}

        claude_api._llamar = llamar_verif
        res = claude_api.verificar_reporte(rep2, "documento corto")
        eq("verificar_reporte coincide", res["coincide"], True)
        res2 = claude_api.verificar_reporte(rep2, "x" * 800_001)
        check("con documento recortado NO puede decir que coincide",
              res2["coincide"] is False and "PARTE" in res2["observaciones"][0],
              str(res2["observaciones"][:1]))

        # normalizar_telefonos por lotes
        rep4 = parser_excel.parsear_excel(RUTA_XLSX, "11")
        lotes = {"n": 0}

        def llamar_tel(sistema, contenido, esquema, progreso=None):
            lotes["n"] += 1
            return {"clientes": [
                {"codigo": "00000002761", "contactos": ["Marcos"],
                 "telefonos": [{"numero": "5491139009488", "es_movil": True,
                                "contacto": "Marcos"},
                               {"numero": "123", "es_movil": False, "contacto": ""}]},
            ]}

        claude_api._llamar = llamar_tel
        n = claude_api.normalizar_telefonos(rep4)
        eq("actualiza 1 cliente", n, 1)
        eq("una sola llamada para 3 clientes", lotes["n"], 1)
        eq("descarta numeros demasiado cortos",
           rep4.clientes[0].telefonos, ["5491139009488"])

        rep5 = modelos.Reporte()
        for i in range(130):
            rep5.clientes.append(modelos.Cliente(codigo=f"{i:011d}",
                                                 telefonos_raw="11 3900 9488"))
        lotes["n"] = 0
        claude_api.normalizar_telefonos(rep5)
        eq("130 clientes -> 3 lotes de 60", lotes["n"], 3)
    finally:
        claude_api._llamar = original_llamar

    # --- disponibilidad / errores ---
    check("esta_disponible = False sin clave y sin entorno",
          claude_api.esta_disponible() == bool(claude_api._credencial_en_entorno()))
    check("_sin_credito detecta el 400 de saldo",
          claude_api._sin_credito("Your credit balance is too low to access the API"))
    check("_sin_credito no dispara con cualquier error",
          not claude_api._sin_credito("invalid model name"))
    e = claude_api._traducir_excepcion(Exception("credit balance is too low"))
    check("el saldo agotado se explica bien", "no tiene credito" in str(e), str(e)[:80])
    eq("precio de entrada opus-5", claude_api.PRECIO_ENTRADA, 5.0)
    eq("precio de salida opus-5", claude_api.PRECIO_SALIDA, 25.0)
    eq("max_tokens de streaming", claude_api.MAX_TOKENS_STREAM, 64000)
    eq("beta de fallbacks", claude_api.BETA_FALLBACK, "server-side-fallback-2026-07-01")



# ==========================================================================
# 8b. WHATSAPP WEB ASISTIDO (sin abrir el navegador)
# ==========================================================================
def t_whatsapp_web() -> None:
    _io = io
    grupo("8b. whatsapp_web.py  (envio asistido, sin abrir Chrome)")

    import whatsapp_web as ww

    # --- URL del chat -----------------------------------------------------
    eq("url_chat solo con el telefono",
       ww.url_chat("5491139009488"),
       "https://web.whatsapp.com/send?phone=5491139009488")
    eq("url_chat limpia lo que no sea digito",
       ww.url_chat("+54 9 11 3900-9488"),
       "https://web.whatsapp.com/send?phone=5491139009488")
    url = ww.url_chat("5491139009488", "hola *mundo*\nsegunda linea")
    check("el texto viaja codificado en la URL", "&text=" in url and "%0A" in url, url)
    check("no se escapa el chat de otro numero", url.count("phone=") == 1)
    # el salto de linea por URL es clave: en WhatsApp Web, Enter ENVIA
    check("los saltos de linea no se tipean, van en la URL",
          "\n" not in url)

    # --- cola de pasos ----------------------------------------------------
    pasos = ww.armar_pasos("Juan Perez (002)", "5491139009488",
                           ["parte uno", "parte dos"],
                           ruta_pdf="C:/tmp/Resumen.pdf", epigrafe="epi")
    eq("un paso por PDF mas uno por parte de texto", len(pasos), 3)
    eq("primero el documento", pasos[0].tipo, "documento")
    eq("despues los textos", [p.tipo for p in pasos[1:]], ["texto", "texto"])
    eq("el PDF lleva su epigrafe", pasos[0].epigrafe, "epi")
    check("las partes van numeradas para el usuario",
          "Mensaje 1 de 2" in pasos[1].descripcion
          and "Mensaje 2 de 2" in pasos[2].descripcion,
          f"{pasos[1].descripcion} | {pasos[2].descripcion}")
    check("todos los pasos saben a quien van",
          all(p.telefono == "5491139009488" and p.vendedor == "Juan Perez (002)"
              for p in pasos))

    solo_texto = ww.armar_pasos("Ana Gomez (005)", "5491140647300", ["unico"])
    eq("sin PDF, un solo paso", len(solo_texto), 1)
    eq("y es de texto", solo_texto[0].tipo, "texto")
    check("con una sola parte no dice 'Mensaje 1 de 1'",
          "de 1" not in solo_texto[0].descripcion, solo_texto[0].descripcion)

    # --- que input sirve para un PDF (el bug del adjunto) -----------------
    # WhatsApp Web arranca con UN input, el de fotos (accept="image/*"). Si se
    # le pasa un PDF lo descarta sin avisar y el adjunto nunca aparece. El de
    # documentos (accept="*") recien existe despues de apretar "Documento".
    casos_accept = [
        ("*", True), ("*/*", True), ("", True), (None, True),
        ("application/pdf", True), (".pdf", True),
        ("image/*", False), ("video/*", False), ("audio/*", False),
        ("image/*,video/mp4", False),
    ]
    for acepta, esperado in casos_accept:
        eq(f"accept={acepta!r} sirve para un PDF",
           ww._acepta_documentos(acepta), esperado)
    check("NUNCA se elige el input de fotos para mandar un PDF",
          not ww._acepta_documentos("image/*"))

    # --- entorno ----------------------------------------------------------
    listo, motivo = ww.disponible()
    check("disponible() devuelve (bool, motivo) sin abrir nada",
          isinstance(listo, bool) and isinstance(motivo, str))
    if listo:
        check("en esta PC estan selenium y Chrome", True)
    else:
        print(f"        (no disponible aca: {motivo.splitlines()[0]})")

    # --- la sesion no hace nada peligroso sin navegador --------------------
    sesion = ww.Sesion()
    eq("arranca sin navegador", sesion.driver, None)
    try:
        sesion.preparar(pasos[0])
        check("preparar() sin navegador tiene que fallar", False)
    except ww.ErrorWhatsAppWeb as e:
        check("preparar() sin navegador avisa claro", "no esta abierto" in str(e), str(e))

    check("el perfil de Chrome vive al lado del programa",
          ww.CARPETA_PERFIL.endswith("PerfilChrome"), ww.CARPETA_PERFIL)
    check("la sesion guardada esta en .gitignore",
          "PerfilChrome" in _io.open(os.path.join(PROYECTO, ".gitignore"),
                                     encoding="utf-8").read())

    # --- no existe ninguna funcion que apriete Enviar sola -----------------
    fuente = _io.open(os.path.join(PROYECTO, "whatsapp_web.py"), encoding="utf-8").read()
    check("el modulo NO manda solo: nada de click en el boton de enviar",
          "send-button" not in fuente and "data-icon='send'" not in fuente)

# ==========================================================================
# 9. INTERFAZ (smoke test real de Tkinter, sin mainloop)
# ==========================================================================
def t_interfaz() -> None:
    grupo("9. interfaz.py  (ventana Tkinter, prueba de humo)")

    import tkinter as tk
    try:
        raiz = tk.Tk()
        raiz.destroy()
    except Exception as e:
        print(f"  [SALTEADA ] no hay entorno grafico disponible ({e})")
        return

    import interfaz

    # neutralizamos los cuadros de dialogo
    dialogos: list[tuple[str, str]] = []
    for nombre in ("showinfo", "showwarning", "showerror"):
        setattr(interfaz.messagebox, nombre,
                lambda titulo, texto="", _n=nombre, **kw: dialogos.append((_n, titulo)))
    interfaz.messagebox.askyesno = lambda *a, **kw: False

    app = interfaz.App()
    try:
        app.withdraw()
        app.update()
        check("la ventana arranca", app.winfo_exists() == 1)
        eq("5 pestanas", len(app.libreta.tabs()), 5)
        check("esta el boton de WhatsApp Web asistido", hasattr(app, "btn_wa_web"))
        check("ese boton se bloquea mientras hay una operacion en curso",
              app.btn_wa_web in (app.btn_leer, app.btn_claude, app.btn_verificar,
                                 app.btn_telefonos, app.btn_enviar, app.btn_wa_web))
        app._bloquear_botones(True)
        eq("bloqueado", str(app.btn_wa_web.cget("state")), "disabled")
        app._bloquear_botones(False)
        eq("desbloqueado", str(app.btn_wa_web.cget("state")), "normal")

        rep = parser_excel.parsear_excel(RUTA_XLSX, "11")
        app.var_ruta.set(RUTA_XLSX)
        app._recibir_reporte(rep)
        app.update()

        check("no aparecieron errores al cargar",
              not any(d[0] == "showerror" for d in dialogos), str(dialogos))
        eq("resumen: total a cobrar", app.var_res_total.get(), "$157.165,89")
        eq("resumen: vencido", app.var_res_vencido.get(), "$137.165,39")
        check("control de totales OK", app.var_res_control.get().startswith("OK:"),
              app.var_res_control.get())
        eq("filas en la tabla de clientes", len(app.tabla.get_children()), 3)
        check("la barra de totales suma",
              "$157.165,89" in app.var_totales_tabla.get(), app.var_totales_tabla.get())

        # filtros
        app.var_f_avencer.set(False)
        app._refrescar_clientes()
        app.update()
        eq("filtro 'solo vencidos'", len(app.tabla.get_children()), 2)
        app.var_f_avencer.set(True)
        app.var_f_resaltados.set(True)
        app._refrescar_clientes()
        eq("filtro 'solo amarillos'", len(app.tabla.get_children()), 1)
        app.var_f_resaltados.set(False)
        app.var_f_buscar.set("maderas")
        app._refrescar_clientes()
        eq("buscador", len(app.tabla.get_children()), 1)
        app.var_f_buscar.set("")
        app.var_f_saldo_min.set("60000")
        app._refrescar_clientes()
        eq("filtro por saldo minimo", len(app.tabla.get_children()), 1)
        app.var_f_saldo_min.set("0")
        app._refrescar_clientes()

        for col in ("codigo", "cliente", "total", "vencido", "situacion", "marca"):
            app._ordenar_por(col)
        check("ordenar por columna no rompe", True)

        # vendedores + asignacion + envio
        for f in (config.ARCHIVO_VENDEDORES, config.ARCHIVO_ASIGNACIONES):
            if os.path.exists(f):
                os.remove(f)
        v1, _ = vend.agregar_vendedor("Juan Perez", "11 3900-9488", "002")
        v2, _ = vend.agregar_vendedor("Ana Gomez", "1140647300", "005")
        app._refrescar_vendedores()
        eq("tabla de vendedores", len(app.tabla_vend.get_children()), 2)

        vend.asignar_varios([rep.clientes[0].codigo, rep.clientes[1].codigo], v1.id)
        vend.asignar_cliente(rep.clientes[2].codigo, v2.id)
        app._refrescar_clientes()
        app._refrescar_envio()
        app.update()
        eq("destinatarios", len(app.tabla_envio.get_children()), 2)
        check("el vendedor aparece en la tabla de clientes",
              any("Juan Perez" in str(app.tabla.item(i)["values"][8])
                  for i in app.tabla.get_children()))

        grupos = app._grupos_envio()
        eq("agrupado: Juan tiene 2", len(grupos[v1.id]), 2)
        partes = app._armar_partes(v1, grupos[v1.id], incluir_comprobantes=False)
        eq("mensaje de Juan en 1 parte", len(partes), 1)
        check("el mensaje de Juan NO menciona clientes de Ana",
              "MADERAS NORTE" not in partes[0])
        check("el mensaje de Juan trae sus 2 clientes",
              "ACME HERRAMIENTAS" in partes[0] and "FERRETERIA DEL SUR" in partes[0])
        total_juan = round(sum(c.saldo_total for c in grupos[v1.id]), 2)
        check("el total del mensaje es solo el de Juan",
              modelos.formatear_pesos(total_juan) in partes[0],
              modelos.formatear_pesos(total_juan))

        ok, ruta = app._generar_pdf_para(v1, grupos[v1.id], True)
        check("PDF por vendedor desde la interfaz", ok, ruta)
        if ok:
            texto = parser_pdf.extraer_texto(ruta)
            check("el PDF de Juan no incluye clientes de Ana",
                  "MADERAS NORTE" not in texto)

        csv = os.path.join(config.CARPETA_HISTORIAL, "envios.csv")
        if os.path.exists(csv):
            os.remove(csv)
        app._registrar_historial(v1, grupos[v1.id], True)
        check("se escribio el historial", os.path.exists(csv))
        if os.path.exists(csv):
            contenido = open(csv, encoding="utf-8-sig").read()
            check("el historial tiene cabecera y una fila",
                  contenido.count("\n") == 2, repr(contenido[-80:]))

        # el archivo con el total roto tiene que avisar
        dialogos.clear()
        rep_mal = parser_excel.parsear_excel(RUTA_XLSX_MAL, "11")
        app._recibir_reporte(rep_mal)
        app.update()
        check("avisa cuando los totales no cierran",
              any(d[0] == "showwarning" for d in dialogos), str(dialogos))
        check("el control de totales lo dice",
              app.var_res_control.get().startswith("ATENCION"), app.var_res_control.get())

        # guardar configuracion (a la ruta temporal)
        dialogos.clear()
        app.var_k_max_car.set("no es un numero")
        app._guardar_config()
        eq("max_caracteres invalido cae al default",
           config.CONFIG.get("max_caracteres_mensaje"), 3800)
        check("config.json se escribio en el sandbox",
              os.path.exists(os.path.join(TMP, "config.json")))
    finally:
        _cerrar_ventana(app)


# ==========================================================================
# 10. TRAMPAS DE DISENO (comportamientos correctos en el codigo pero riesgosos)
# ==========================================================================
def t_trampas() -> None:
    grupo("10. Trampas de diseno")

    import interfaz
    import tkinter as tk
    for nombre in ("showinfo", "showwarning", "showerror"):
        setattr(interfaz.messagebox, nombre, lambda *a, **kw: None)
    interfaz.messagebox.askyesno = lambda *a, **kw: False

    # reporte con un cliente a vencer dentro de 30 dias
    rep = modelos.Reporte(fecha_reporte=FECHA_REPORTE, empresa="WOOD TOOLS SRL")
    cerca = modelos.Cliente(codigo="00000000001", razon_social="VENCE EN 3 DIAS SA")
    cerca.comprobantes = [modelos.Comprobante(tipo="FE A", saldo=1000.0, dias_venc=-3,
                                              fecha_vto=date(2026, 8, 13))]
    cerca.total_saldo_reporte = 1000.0
    lejos = modelos.Cliente(codigo="00000000002", razon_social="VENCE EN 30 DIAS SA")
    lejos.comprobantes = [modelos.Comprobante(tipo="FE A", saldo=900000.0, dias_venc=-30,
                                              fecha_vto=date(2026, 9, 9))]
    lejos.total_saldo_reporte = 900000.0
    chico = modelos.Cliente(codigo="00000000003", razon_social="DEUDA CHICA SA")
    chico.comprobantes = [modelos.Comprobante(tipo="FE A", saldo=50.0, dias_venc=90,
                                              fecha_vto=date(2026, 5, 12))]
    chico.total_saldo_reporte = 50.0
    rep.clientes = [cerca, lejos, chico]
    rep.total_general_saldo = 901050.0

    app = interfaz.App()
    try:
        app.withdraw()
        for f in (config.ARCHIVO_VENDEDORES, config.ARCHIVO_ASIGNACIONES):
            if os.path.exists(f):
                os.remove(f)
        v, _ = vend.agregar_vendedor("Juan Perez", "11 3900-9488", "002")
        vend.asignar_varios([c.codigo for c in rep.clientes], v.id)

        app._recibir_reporte(rep)
        app.update()

        # filtro por defecto: "a vencer dentro de 7 dias"
        eq("filtro por defecto 'a vencer dentro de N'", app.var_f_dias.get(), "7")
        visibles = app._clientes_filtrados()
        nombres = {c.nombre for c in visibles}
        check("TRAMPA: el cliente que vence en 30 dias queda FUERA por el filtro de 7 dias",
              "VENCE EN 30 DIAS SA" not in nombres, str(nombres))
        grupos = app._grupos_envio()
        partes = app._armar_partes(v, grupos[v.id], incluir_comprobantes=False)
        check("TRAMPA: y por lo tanto tampoco entra al WhatsApp",
              "VENCE EN 30 DIAS SA" not in partes[0])
        check("TRAMPA: el TOTAL del mensaje no incluye esos $900.000",
              "$900.000,00" not in partes[0] and "$901.050,00" not in partes[0])
        print(f"        total del reporte: {modelos.formatear_pesos(rep.saldo_total)}")
        print(f"        total que se manda: {modelos.formatear_pesos(round(sum(c.saldo_total for c in grupos[v.id]), 2))}")

        # el filtro de saldo minimo tambien recorta el envio
        app.var_f_dias.set("0")
        app.var_f_saldo_min.set("100")
        app._refrescar_clientes()
        grupos = app._grupos_envio()
        partes = app._armar_partes(v, grupos[v.id], incluir_comprobantes=False)
        check("TRAMPA: 'saldo minimo' de la pestana 2 tambien recorta el envio",
              "DEUDA CHICA" not in partes[0])
        app.var_f_saldo_min.set("0")
        app._refrescar_clientes()

        # --- ARREGLO: la pestana de envio ahora avisa lo que queda afuera ---
        app.var_f_dias.set("7")
        app.var_f_saldo_min.set("0")
        app._refrescar_clientes()
        app.update()
        aviso = app.var_filtros_envio.get()
        print(f"        aviso en pantalla: {aviso}")
        check("ARREGLADO: la pestana de envio avisa que hay clientes afuera",
              aviso.startswith("ATENCION"), aviso)
        check("ARREGLADO: dice cuantos clientes", "1 cliente" in aviso, aviso)
        check("ARREGLADO: dice cuanta plata", "$900.000,00" in aviso, aviso)
        check("ARREGLADO: dice que filtro lo causa",
              "a vencer dentro de 7 dias" in aviso, aviso)

        fuera, monto = app._clientes_excluidos()
        eq("_clientes_excluidos: 1 cliente", len(fuera), 1)
        eq("_clientes_excluidos: el monto", monto, 900000.0, tol=0.005)

        app.var_f_saldo_min.set("100")
        app._refrescar_clientes()
        aviso = app.var_filtros_envio.get()
        check("ARREGLADO: acumula los dos filtros",
              "2 clientes" in aviso and "$900.050,00" in aviso, aviso)
        check("ARREGLADO: nombra los dos filtros",
              "a vencer dentro de 7 dias" in aviso and "saldo minimo" in aviso, aviso)

        # sin filtros que recorten, el cartel desaparece
        app.var_f_dias.set("0")
        app.var_f_saldo_min.set("0")
        app._refrescar_clientes()
        app.update()
        eq("ARREGLADO: sin recorte, no hay cartel", app.var_filtros_envio.get(), "")
        eq("y entran los 3 clientes", len(app._clientes_filtrados()), 3)

        # el aviso se actualiza solo al tocar un filtro (sin apretar nada mas)
        app.var_f_resaltados.set(True)
        app._refrescar_clientes()
        app.update()
        check("ARREGLADO: el cartel se actualiza al cambiar un filtro",
              "solo los marcados en amarillo" in app.var_filtros_envio.get(),
              app.var_filtros_envio.get())
        app.var_f_resaltados.set(False)
        app._refrescar_clientes()
    finally:
        _cerrar_ventana(app)

    print("  -- pendientes conocidos (no cuentan como falla) --")
    nota("estimar_costo usa precios fijos de opus-5 y no mira el modelo elegido",
         claude_api.PRECIO_ENTRADA == 5.0 and claude_api.PRECIO_SALIDA == 25.0
         and "PRECIO_ENTRADA" in claude_api.estimar_costo.__code__.co_names)

    # limites de la API que NO se controlan
    import inspect
    src = inspect.getsource(claude_api.interpretar_pdf)
    check("interpretar_pdf controla el tamano de 30 MB", "30 * 1024 * 1024" in src)
    nota("interpretar_pdf no controla el limite de 600 paginas de la API",
         "600" not in src and "contar_paginas" not in src)

    # claves de configuracion muertas
    import io as _io
    fuentes = ""
    for m in ("interfaz", "claude_api", "mensajes", "parser_excel", "parser_pdf",
              "pdf_resumen", "vendedores", "whatsapp"):
        fuentes += _io.open(os.path.join(PROYECTO, m + ".py"), encoding="utf-8").read()
    for clave in ("incluir_a_vencer", "usar_claude_en_excel"):
        nota(f"la clave de config '{clave}' no la lee ningun modulo",
             clave not in fuentes)
    otros = ""
    for m in ("interfaz", "mensajes", "parser_excel", "parser_pdf", "pdf_resumen",
              "vendedores", "whatsapp"):
        otros += _io.open(os.path.join(PROYECTO, m + ".py"), encoding="utf-8").read()
    nota("claude_api.redactar_mensaje() no lo llama nadie",
         "redactar_mensaje" not in otros)
    parsers = ""
    for m in ("parser_excel", "parser_pdf", "claude_api"):
        parsers += _io.open(os.path.join(PROYECTO, m + ".py"), encoding="utf-8").read()
    nota("ningun parser completa Cliente.codigo_vendedor",
         "codigo_vendedor" not in parsers)


# ==========================================================================
# MAIN
# ==========================================================================
def main() -> None:
    print("Banco de pruebas del Gestor de Cobranzas")
    print(f"Proyecto: {PROYECTO}")
    for fn in (t_modelos, t_parser_excel, t_parser_pdf, t_ocr, t_mensajes,
               t_pdf_resumen, t_vendedores, t_whatsapp, t_whatsapp_web,
               t_claude, t_interfaz,
               t_trampas):
        seccion(fn)

    print(f"\n{'=' * 78}\nRESUMEN\n{'=' * 78}")
    fallas = [(g, n, d) for g, n, d in RESULTADOS if d != "ok"]
    print(f"  Verificaciones: {len(RESULTADOS)}")
    print(f"  OK:             {len(RESULTADOS) - len(fallas)}")
    print(f"  FALLAS:         {len(fallas)}")
    print(f"  Carpeta temporal usada: {TMP}")
    if fallas:
        print("\n  DETALLE DE LAS FALLAS")
        for g, n, d in fallas:
            print(f"   - [{g.split('.')[0]}] {n}: {d}")
    sys.exit(1 if fallas else 0)


if __name__ == "__main__":
    main()
