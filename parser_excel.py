# -*- coding: utf-8 -*-
"""
Parser deterministico del reporte "A cobrar por cliente" en Excel.

Soporta:
  - .xls  (formato viejo BIFF, el que exporta Gestion Comercial) via xlrd
  - .xlsx / .xlsm                                                via openpyxl

Lee valores Y colores de relleno, para poder detectar lo marcado en amarillo.
No necesita tener Excel instalado.

Layout detectado en "LISTADO VENCIDO.XLS" (columnas 0-based):
  0=A talonario / codigo cliente   1=B tipo comprobante   2=C nro / razon social
  3=D condicion de pago            4=E fecha emision      5=F fecha vto / localidad
  6=G nro de cuota / telefonos     7=H importe original   8=I saldo / contacto
  9=J dias vencidos               10=K FC credito
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Iterable

from modelos import (
    Cliente,
    Comprobante,
    Reporte,
    telefonos_de_texto,
    sin_acentos,
)

# --------------------------------------------------------------------------
# INDICES DE COLUMNA
# --------------------------------------------------------------------------
COL_A, COL_B, COL_C, COL_D, COL_E = 0, 1, 2, 3, 4
COL_F, COL_G, COL_H, COL_I, COL_J, COL_K = 5, 6, 7, 8, 9, 10

# Alias legibles
C_TALONARIO = COL_A
C_TIPO = COL_B
C_NUMERO = COL_C
C_CONDICION = COL_D
C_EMISION = COL_E
C_VTO = COL_F
C_CUOTA = COL_G
C_ORIGINAL = COL_H
C_SALDO = COL_I
C_DIAS = COL_J
C_FCCRED = COL_K

C_COD_CLIENTE = COL_A
C_RAZON = COL_C
C_LOCALIDAD = COL_F
C_TELEFONOS = COL_G
C_CONTACTO = COL_I


# --------------------------------------------------------------------------
# GRILLA NORMALIZADA (independiente de xlrd / openpyxl)
# --------------------------------------------------------------------------
@dataclass
class Celda:
    valor: Any = None                       # str | float | date | None
    color: tuple[int, int, int] | None = None   # relleno solido, si tiene

    @property
    def texto(self) -> str:
        v = self.valor
        if v is None:
            return ""
        if isinstance(v, str):
            return v.strip()
        if isinstance(v, (date, datetime)):
            return v.strftime("%d/%m/%Y")
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()

    @property
    def es_numero(self) -> bool:
        return isinstance(self.valor, (int, float)) and not isinstance(self.valor, bool)

    @property
    def numero(self) -> float | None:
        if self.es_numero:
            return float(self.valor)
        # numeros escritos como texto: "37.165,39" o "37165.39"
        t = self.texto
        if not t:
            return None
        m = re.fullmatch(r"-?[\d\.]*\d(?:,\d+)?|-?\d+(?:\.\d+)?", t.replace(" ", ""))
        if not m:
            return None
        limpio = t.replace(" ", "")
        if "," in limpio:
            limpio = limpio.replace(".", "").replace(",", ".")
        try:
            return float(limpio)
        except ValueError:
            return None

    @property
    def fecha(self) -> date | None:
        v = self.valor
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        t = self.texto
        m = re.fullmatch(r"(\d{1,2})[/\-](\d{1,2})[/\-](\d{2,4})", t)
        if m:
            d, mes, a = (int(x) for x in m.groups())
            if a < 100:
                a += 2000
            try:
                return date(a, mes, d)
            except ValueError:
                return None
        return None

    @property
    def resaltado(self) -> bool:
        """Tiene relleno de color distinto de blanco/negro."""
        if self.color is None:
            return False
        r, g, b = self.color
        if (r, g, b) == (255, 255, 255) or (r, g, b) == (0, 0, 0):
            return False
        return True

    @property
    def es_amarillo(self) -> bool:
        if self.color is None:
            return False
        r, g, b = self.color
        return r >= 200 and g >= 180 and b <= 140


VACIA = Celda()


class Grilla:
    """Matriz de celdas 0-based con acceso seguro fuera de rango."""

    def __init__(self, filas: list[list[Celda]], nombre_hoja: str = ""):
        self.filas = filas
        self.nombre_hoja = nombre_hoja

    def __len__(self) -> int:
        return len(self.filas)

    def celda(self, f: int, c: int) -> Celda:
        if 0 <= f < len(self.filas):
            fila = self.filas[f]
            if 0 <= c < len(fila):
                return fila[c]
        return VACIA

    def fila(self, f: int) -> list[Celda]:
        if 0 <= f < len(self.filas):
            return self.filas[f]
        return []

    def fila_vacia(self, f: int) -> bool:
        return all(not c.texto for c in self.fila(f))

    def texto_fila(self, f: int) -> str:
        return " ".join(c.texto for c in self.fila(f) if c.texto).strip()


# --------------------------------------------------------------------------
# LECTORES
# --------------------------------------------------------------------------
def _leer_xls(ruta: str) -> list[Grilla]:
    """Lee .xls con xlrd, incluyendo colores de relleno."""
    import xlrd

    try:
        wb = xlrd.open_workbook(ruta, formatting_info=True)
        con_formato = True
    except Exception:
        # Algunos .xls no traen la tabla de formatos; leemos solo valores.
        wb = xlrd.open_workbook(ruta)
        con_formato = False

    grillas: list[Grilla] = []
    for hoja in wb.sheets():
        filas: list[list[Celda]] = []
        for r in range(hoja.nrows):
            fila: list[Celda] = []
            for c in range(hoja.ncols):
                cel = hoja.cell(r, c)
                valor: Any = None
                if cel.ctype == xlrd.XL_CELL_TEXT:
                    valor = cel.value
                elif cel.ctype == xlrd.XL_CELL_NUMBER:
                    valor = cel.value
                elif cel.ctype == xlrd.XL_CELL_DATE:
                    try:
                        valor = xlrd.xldate_as_datetime(cel.value, wb.datemode)
                    except Exception:
                        valor = cel.value
                elif cel.ctype == xlrd.XL_CELL_BOOLEAN:
                    valor = bool(cel.value)

                color = None
                if con_formato:
                    try:
                        xf = wb.xf_list[cel.xf_index]
                        fondo = xf.background
                        if fondo.fill_pattern == 1:  # solido
                            rgb = wb.colour_map.get(fondo.pattern_colour_index)
                            if rgb:
                                color = tuple(int(x) for x in rgb[:3])  # type: ignore[assignment]
                    except Exception:
                        color = None
                fila.append(Celda(valor, color))
            filas.append(fila)
        grillas.append(Grilla(filas, hoja.name))
    return grillas


def _hex_a_rgb(valor: Any) -> tuple[int, int, int] | None:
    if not isinstance(valor, str):
        return None
    h = valor.strip().upper()
    if h.startswith("#"):
        h = h[1:]
    if len(h) == 8:      # AARRGGBB
        h = h[2:]
    if len(h) != 6 or not re.fullmatch(r"[0-9A-F]{6}", h):
        return None
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _leer_xlsx(ruta: str) -> list[Grilla]:
    """Lee .xlsx/.xlsm con openpyxl, incluyendo colores de relleno."""
    from openpyxl import load_workbook

    wb = load_workbook(ruta, data_only=True)
    grillas: list[Grilla] = []
    for hoja in wb.worksheets:
        filas: list[list[Celda]] = []
        for fila_ws in hoja.iter_rows():
            fila: list[Celda] = []
            for cel in fila_ws:
                color = None
                try:
                    relleno = cel.fill
                    if relleno is not None and relleno.patternType == "solid":
                        color = _hex_a_rgb(getattr(relleno.fgColor, "rgb", None))
                except Exception:
                    color = None
                fila.append(Celda(cel.value, color))
            filas.append(fila)
        grillas.append(Grilla(filas, hoja.title))
    try:
        wb.close()
    except Exception:
        pass
    return grillas


def leer_grillas(ruta: str) -> list[Grilla]:
    ext = os.path.splitext(ruta)[1].lower()
    if ext == ".xls":
        return _leer_xls(ruta)
    if ext in (".xlsx", ".xlsm", ".xltx"):
        return _leer_xlsx(ruta)
    raise ValueError(f"Extension no soportada por el parser de Excel: {ext}")


# --------------------------------------------------------------------------
# CLASIFICACION DE FILAS
# --------------------------------------------------------------------------
_RE_COD_CLIENTE = re.compile(r"^\d{6,}$")
_RE_LOCALIDAD = re.compile(r"^\(\s*\d{3,5}\s*\)")
_RE_TOTAL_CLIENTE = re.compile(
    r"total\s+cliente\s+(\d+)\s*(.*)$", re.IGNORECASE
)

# Textos de cabecera/pie que se repiten en cada hoja del reporte y hay que ignorar
_PREFIJOS_RUIDO = (
    "GESTION COMERCIAL",
    "A COBRAR POR CLIENTE",
    "CLIENTE:",
    "CLIENTE COMPRADOR:",
    "COND.VTA",
    "EXCLUYE",
    "ORDENADO POR",
    "FECHA:",
    "HOJA:",
    "TOTAL GENERAL",
    "CONTINUA EN LA HOJA",
    "SUBTOTAL HOJA",
)
_EXACTOS_RUIDO = {
    "CLIENTE",
    "COMPROBANTE",
    "CONDICION DE PAGO",
    "FECHA",
    "EMISION",
    "VTO.",
    "SALDO",
    "IMPORTE",
    "ORIGINAL",
    "DIAS",
    "VENC.",
    "N DE",
    "NO DE",
    "CUOTA",
    "FC CREDITO",
    "PAGINA",
}


def _norm(txt: str) -> str:
    return sin_acentos(txt or "").upper().strip()


def _es_ruido_cabecera(fila: list[Celda], empresa: str = "") -> bool:
    textos = [_norm(c.texto) for c in fila if c.texto]
    if not textos:
        return False
    if empresa and textos[0] == _norm(empresa):
        return True
    for t in textos:
        for pref in _PREFIJOS_RUIDO:
            if t.startswith(pref):
                return True
    # Fila de titulos de columnas: 2 o mas celdas que son etiquetas exactas
    coincidencias = sum(1 for t in textos if t.replace("º", "").replace("°", "") in _EXACTOS_RUIDO)
    if coincidencias >= 2:
        return True
    # Fila con una sola celda que es exactamente una etiqueta
    if len(textos) == 1 and textos[0] in _EXACTOS_RUIDO:
        return True
    return False


def _es_cabecera_cliente(fila: list[Celda]) -> bool:
    """col A = codigo largo de digitos, col C = razon social, col B vacia."""
    a = fila[C_COD_CLIENTE] if len(fila) > C_COD_CLIENTE else VACIA
    b = fila[C_TIPO] if len(fila) > C_TIPO else VACIA
    c = fila[C_RAZON] if len(fila) > C_RAZON else VACIA
    if b.texto:
        return False
    cod = re.sub(r"\D", "", a.texto)
    if not cod or not _RE_COD_CLIENTE.match(cod):
        return False
    return bool(c.texto) and not c.es_numero


def _es_comprobante(fila: list[Celda]) -> bool:
    """col B = tipo (texto), col H = importe numerico, col A = talonario corto."""
    b = fila[C_TIPO] if len(fila) > C_TIPO else VACIA
    h = fila[C_ORIGINAL] if len(fila) > C_ORIGINAL else VACIA
    if not b.texto:
        return False
    if _norm(b.texto) in _EXACTOS_RUIDO:
        return False
    if _norm(b.texto).startswith("TOTAL"):
        return False
    return h.numero is not None


def _es_total_cliente(fila: list[Celda]) -> re.Match[str] | None:
    for idx in (C_TIPO, C_COD_CLIENTE, C_RAZON, C_CONDICION):
        if len(fila) > idx:
            m = _RE_TOTAL_CLIENTE.search(sin_acentos(fila[idx].texto))
            if m:
                return m
    return None


def _es_total_general(fila: list[Celda]) -> bool:
    for c in fila:
        if _norm(c.texto).startswith("TOTAL GENERAL"):
            return True
    return False


# --------------------------------------------------------------------------
# EXTRACCION DE LA CABECERA DEL REPORTE
# --------------------------------------------------------------------------
_RE_FILTRO_VEND = re.compile(r"Vendedor\s*:\s*(\d+)\s*-\s*(\d+)", re.IGNORECASE)
_RE_FILTRO_COBR = re.compile(r"Cobrador\s*:\s*(\d+)\s*-\s*(\d+)", re.IGNORECASE)


def _extraer_cabecera(grilla: Grilla, reporte: Reporte) -> None:
    """Empresa, titulo, fecha del reporte y filtros aplicados."""
    limite = min(len(grilla), 20)
    for f in range(limite):
        fila = grilla.fila(f)
        a = grilla.celda(f, COL_A).texto
        if f == 0 and a and not reporte.empresa:
            reporte.empresa = a
        if a and not reporte.titulo and _norm(a).startswith("A COBRAR"):
            reporte.titulo = a
        # "Fecha:" en alguna columna, con la fecha en la siguiente
        for c, cel in enumerate(fila):
            if _norm(cel.texto) == "FECHA:" and reporte.fecha_reporte is None:
                siguiente = grilla.celda(f, c + 1)
                fecha = siguiente.fecha
                if fecha is None and siguiente.es_numero:
                    fecha = _serial_a_fecha(siguiente.numero)
                if fecha:
                    reporte.fecha_reporte = fecha
        # Linea de filtros
        if a and ("Vendedor" in a or "Cobrador" in a or a.startswith("Cliente:")):
            reporte.filtros = (reporte.filtros + " " + a).strip()
        if a and (_norm(a).startswith("EXCLUYE") or _norm(a).startswith("ORDENADO")):
            reporte.filtros = (reporte.filtros + " | " + a).strip()

    m = _RE_FILTRO_VEND.search(reporte.filtros or "")
    if m:
        reporte.filtro_vendedor_desde, reporte.filtro_vendedor_hasta = m.group(1), m.group(2)
    m = _RE_FILTRO_COBR.search(reporte.filtros or "")
    if m:
        reporte.filtro_cobrador = m.group(1)


def _serial_a_fecha(serial: float | None) -> date | None:
    """Convierte un serial de Excel (base 1900) a fecha."""
    if serial is None:
        return None
    try:
        from datetime import timedelta

        base = date(1899, 12, 30)  # base 1900 con el bug del 29/02/1900
        return base + timedelta(days=int(serial))
    except Exception:
        return None


# --------------------------------------------------------------------------
# PARSER PRINCIPAL
# --------------------------------------------------------------------------
class _EstadoCliente:
    """Acumula las filas de informacion que siguen a la cabecera de un cliente."""

    def __init__(self, cliente: Cliente):
        self.cliente = cliente
        self.textos_c: list[str] = []
        self.textos_g: list[str] = []
        self.textos_i: list[str] = []
        self.localidades: list[str] = []

    def agregar_fila_info(self, fila: list[Celda]) -> None:
        c = fila[C_RAZON] if len(fila) > C_RAZON else VACIA
        g = fila[C_TELEFONOS] if len(fila) > C_TELEFONOS else VACIA
        i = fila[C_CONTACTO] if len(fila) > C_CONTACTO else VACIA
        f = fila[C_LOCALIDAD] if len(fila) > C_LOCALIDAD else VACIA

        if c.texto and not (c.es_numero and (c.numero or 0) == 0):
            self.textos_c.append(c.texto)
        if g.texto:
            self.textos_g.append(g.texto)
        if i.texto and not i.es_numero:
            self.textos_i.append(i.texto)
        if f.texto and _RE_LOCALIDAD.match(f.texto):
            self.localidades.append(f.texto)

    def cerrar(self, area_defecto: str = "11") -> None:
        cl = self.cliente
        # La ultima linea de la col C es la direccion; las anteriores, nombre de fantasia
        if self.textos_c:
            cl.direccion = self.textos_c[-1]
            if len(self.textos_c) > 1:
                cl.nombre_fantasia = " / ".join(self.textos_c[:-1])
        if self.localidades and not cl.localidad:
            cl.localidad = self.localidades[0]
        if self.textos_i and not cl.contacto:
            cl.contacto = " / ".join(dict.fromkeys(self.textos_i))
        crudos = ([cl.telefonos_raw] if cl.telefonos_raw else []) + self.textos_g
        cl.telefonos_raw = " / ".join(dict.fromkeys([t for t in crudos if t]))
        cl.telefonos = telefonos_de_texto(cl.telefonos_raw, area_defecto)


def parsear_excel(ruta: str, area_telefono_defecto: str = "11") -> Reporte:
    """Lee el reporte de Excel y devuelve un Reporte poblado.

    No lanza excepciones por filas raras: las anota en `reporte.avisos`.
    """
    grillas = leer_grillas(ruta)
    reporte = Reporte(origen_archivo=ruta, motor="excel")

    if not grillas:
        reporte.avisos.append("El archivo no tiene hojas legibles.")
        return reporte

    for grilla in grillas:
        _extraer_cabecera(grilla, reporte)
        _parsear_grilla(grilla, reporte, area_telefono_defecto)

    if not reporte.clientes:
        reporte.avisos.append(
            "No se reconocio ningun cliente. Puede que el formato del reporte haya "
            "cambiado: proba interpretarlo con Claude."
        )
    return reporte


def _parsear_grilla(grilla: Grilla, reporte: Reporte, area_defecto: str) -> None:
    estado: _EstadoCliente | None = None
    vio_comprobante = False

    def cerrar(motivo: str = "") -> None:
        nonlocal estado, vio_comprobante
        if estado is not None:
            estado.cerrar(area_defecto)
            reporte.clientes.append(estado.cliente)
            if motivo:
                reporte.avisos.append(motivo)
        estado = None
        vio_comprobante = False

    for f in range(len(grilla)):
        fila = grilla.fila(f)
        nro_fila = f + 1  # para mostrarle al usuario la fila real del Excel

        if grilla.fila_vacia(f):
            continue

        # --- total general: cierra todo ---
        if _es_total_general(fila):
            cerrar()
            h = grilla.celda(f, C_ORIGINAL).numero
            i = grilla.celda(f, C_SALDO).numero
            if h is not None:
                reporte.total_general_original = h
            if i is not None:
                reporte.total_general_saldo = i
            continue

        # --- total por cliente: cierra el cliente actual ---
        m_total = _es_total_cliente(fila)
        if m_total is not None:
            h = grilla.celda(f, C_ORIGINAL).numero
            i = grilla.celda(f, C_SALDO).numero
            resaltado_fila = any(c.es_amarillo for c in fila)
            if estado is None:
                # Total sin cabecera previa: reconstruimos desde el texto
                codigo = m_total.group(1)
                nombre = (m_total.group(2) or "").strip()
                cli = Cliente(codigo=codigo, razon_social=nombre, fila_origen=nro_fila)
                estado = _EstadoCliente(cli)
                reporte.avisos.append(
                    f"Fila {nro_fila}: 'Total cliente' sin cabecera de cliente; "
                    f"se reconstruyo desde el texto."
                )
            estado.cliente.total_original_reporte = h
            estado.cliente.total_saldo_reporte = i
            if resaltado_fila:
                estado.cliente.resaltado = True
            cerrar()
            continue

        # --- cabecera de cliente ---
        if _es_cabecera_cliente(fila):
            cerrar()
            cod = re.sub(r"\D", "", grilla.celda(f, C_COD_CLIENTE).texto)
            cli = Cliente(
                codigo=cod,
                razon_social=grilla.celda(f, C_RAZON).texto,
                telefonos_raw=grilla.celda(f, C_TELEFONOS).texto,
                fila_origen=nro_fila,
            )
            loc = grilla.celda(f, C_LOCALIDAD).texto
            if loc and _RE_LOCALIDAD.match(loc):
                cli.localidad = loc
            contacto = grilla.celda(f, C_CONTACTO)
            if contacto.texto and not contacto.es_numero:
                cli.contacto = contacto.texto
            cli.resaltado = any(c.es_amarillo for c in fila)
            estado = _EstadoCliente(cli)
            vio_comprobante = False
            continue

        # --- comprobante ---
        if _es_comprobante(fila):
            if estado is None:
                reporte.avisos.append(
                    f"Fila {nro_fila}: comprobante sin cliente asociado; se ignoro."
                )
                continue
            comp = _armar_comprobante(grilla, f)
            estado.cliente.comprobantes.append(comp)
            if comp.resaltado:
                estado.cliente.resaltado = True
            vio_comprobante = True
            continue

        # --- ruido de cabecera de hoja ---
        if _es_ruido_cabecera(fila, reporte.empresa):
            continue

        # --- lineas de informacion del cliente (fantasia, direccion, localidad) ---
        if estado is not None and not vio_comprobante:
            estado.agregar_fila_info(fila)
            if any(c.es_amarillo for c in fila):
                estado.cliente.resaltado = True
            continue

        # cualquier otra cosa se ignora en silencio (lineas de separacion, etc.)

    cerrar()


def _armar_comprobante(grilla: Grilla, f: int) -> Comprobante:
    def cel(c: int) -> Celda:
        return grilla.celda(f, c)

    emision = cel(C_EMISION).fecha or _serial_a_fecha(cel(C_EMISION).numero)
    vto = cel(C_VTO).fecha or _serial_a_fecha(cel(C_VTO).numero)

    dias_val = cel(C_DIAS).numero
    dias = int(round(dias_val)) if dias_val is not None else None

    original = cel(C_ORIGINAL).numero or 0.0
    saldo_val = cel(C_SALDO).numero
    saldo = saldo_val if saldo_val is not None else original

    comp = Comprobante(
        talonario=cel(C_TALONARIO).texto,
        tipo=cel(C_TIPO).texto,
        numero=cel(C_NUMERO).texto,
        condicion_pago=cel(C_CONDICION).texto,
        fecha_emision=emision,
        fecha_vto=vto,
        cuota=cel(C_CUOTA).texto,
        importe_original=round(float(original), 2),
        saldo=round(float(saldo), 2),
        dias_venc=dias,
        fc_credito=cel(C_FCCRED).texto,
        resaltado=any(c.es_amarillo for c in grilla.fila(f)),
        fila_origen=f + 1,
    )
    return comp


# --------------------------------------------------------------------------
# RECALCULO DE DIAS CONTRA UNA FECHA DE REFERENCIA
# --------------------------------------------------------------------------
def recalcular_dias(reporte: Reporte, referencia: date) -> int:
    """Recalcula `dias_venc` de cada comprobante contra `referencia`.

    Devuelve la cantidad de comprobantes modificados. Se usa cuando el reporte
    es de otro dia y se quiere clasificar vencido/a vencer contra HOY.
    """
    cambios = 0
    for cli in reporte.clientes:
        for comp in cli.comprobantes:
            if comp.fecha_vto is None:
                continue
            nuevos = (referencia - comp.fecha_vto).days
            if comp.dias_venc != nuevos:
                comp.dias_venc = nuevos
                cambios += 1
    return cambios


# --------------------------------------------------------------------------
# VOLCADO EN TEXTO (lo que se le manda a Claude cuando hay que interpretar)
# --------------------------------------------------------------------------
def volcar_para_claude(ruta: str, max_filas: int = 4000) -> str:
    """Serializa el Excel a texto compacto, anotando el resaltado amarillo.

    Formato por fila:
        F19 | A=71 | B=FE A | C=00072065 | ... | I=37165.39 [AMARILLO] | J=10
    """
    grillas = leer_grillas(ruta)
    lineas: list[str] = []
    for grilla in grillas:
        lineas.append(f"### HOJA: {grilla.nombre_hoja}")
        for f in range(min(len(grilla), max_filas)):
            partes: list[str] = []
            for c, cel in enumerate(grilla.fila(f)):
                if not cel.texto and not cel.resaltado:
                    continue
                letra = chr(ord("A") + c) if c < 26 else f"c{c}"
                marca = " [AMARILLO]" if cel.es_amarillo else (
                    " [RESALTADO]" if cel.resaltado else ""
                )
                partes.append(f"{letra}={cel.texto}{marca}")
            if partes:
                lineas.append(f"F{f + 1} | " + " | ".join(partes))
    return "\n".join(lineas)
